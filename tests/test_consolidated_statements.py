"""Behavioral contract tests: no production credentials, database, or ERP access."""
import copy
import importlib.util
import sys
import unittest
from decimal import Decimal
from pathlib import Path

SCRIPTS=Path(__file__).resolve().parents[1]/"skills/consolidated-statements/scripts"
sys.path.insert(0,str(SCRIPTS))
import engine

class StatementContractTests(unittest.TestCase):
    def sheet(self,period="2026年02月"):
        labels=engine.CATALOG["small"]["is"]
        rows=[["利润表","","",""],["甲骨易（北京）语言科技股份有限公司山东分公司",period,"单位：元",""],["项目","行次","本年累计金额","本月金额"]]
        for key,label in sorted(labels.items(),key=lambda p:int(p[0])):
            rows.append([label,key,"",""])
        for sid,year,month in [(2,200,100),(14,70,30),(21,-270,-130),(30,-270,-130),(32,-270,-130)]:
            rows[sid+2][2:]=[year,month]
        return rows

    def parse(self,rows,period="202602"):
        return engine.parse_sheet(rows,"利润表",period,"input.xls","abc")

    def test_distinct_month_and_ytd_and_source_immutability(self):
        source=self.parse(self.sheet());target=engine.classify(source)
        self.assertEqual(source.values[2,0].amount,Decimal(100))
        self.assertEqual(source.values[5,0].amount,Decimal(30))
        self.assertEqual(target.values[2,0].amount,Decimal(130))
        self.assertEqual(target.values[2,1].amount,Decimal(270))
        self.assertEqual(target.values[5,1].amount,Decimal(0))
        self.assertEqual(target.values[24,0].amount,source.values[24,0].amount)

    def test_jinan_legal_name_and_iso_month(self):
        rows=self.sheet()
        rows[1][0]="核算单位：甲骨易(济南)科技有限公司"
        rows[1][1]="2026-02"
        self.assertEqual(self.parse(rows).company,"JINAN")

    def test_known_labels_accept_sichuan_variants_and_reordered_ids(self):
        rows=self.sheet()
        rows[1][0]="甲骨易（北京）语言科技股份有限公司四川分公司"
        rows[12][0]="教育费附加、矿产资源补偿费、排污费"
        rows[14][0]="其中：商品维护费"
        rows[22][0]='加：投资收益（亏损以"-"号填列）'
        rows[34][0]='四、净利润（净亏损以"-"号填列）'
        for row in rows[3:]:
            row[1]=int(row[1])+100
        result=self.parse(rows)
        self.assertEqual(result.company,"SICHUAN")
        self.assertEqual(result.values[24,0].amount,Decimal(-130))

    def test_missing_zero_detail_does_not_drop_entire_report(self):
        rows=self.sheet()
        del rows[6]
        result=self.parse(rows)
        self.assertEqual(result.values[24,0].amount,Decimal(-130))

    def test_missing_report_states_are_distinct(self):
        import tempfile
        from workbook import build_workbook
        with tempfile.TemporaryDirectory() as tmp:
            result=build_workbook([self.parse(self.sheet())],"202602",Path(tmp)/"test.xlsx",input_issues=[
                {"company":"HEAD","kind":"is","status":"rejected","message":"金额列待确认"},
                {"company":"HUNAN_BRANCH","kind":"bs","status":"missing","message":"页面等待失败"},
            ])
        rows={row["company"]:row for row in result["coverage"]}
        self.assertEqual(rows["HEAD"]["is"],"解析失败")
        self.assertEqual(rows["HUNAN_BRANCH"]["bs"],"取数失败")
        self.assertEqual(rows["JINAN"]["bs"],"未上传")
        self.assertFalse(result["coverage_complete"])

    def test_august_ambiguous_current_period_still_rejected(self):
        rows=self.sheet("2026年08月")
        rows[2][2]="本期金额"
        with self.assertRaises(engine.SourceError):
            self.parse(rows,"202608")

    def test_verified_kingdee_cashflow_ytd_profile_is_scoped(self):
        rows=[["现金流量表"],["编制单位："+engine.COMPANIES["HEAD"][1],"2026年08月","单位：元"],["项目","行次","本月金额","本期金额"]]
        for item in engine.CATALOG["standard"]["cf"]["rows"]:
            if item["id"]:
                rows.append([item["label"],item["id"],10,80])
        filename=engine.COMPANIES["HEAD"][1]+"_现金流量表__202608期.xlsx"
        parsed=engine.parse_sheet(rows,"202608","202608",filename,"test")
        self.assertEqual(parsed.values[1,1].amount,Decimal(80))
        self.assertEqual(parsed.metric_basis,"kingdee-standard-cashflow-ytd-v1")
        with self.assertRaises(engine.SourceError):
            engine.parse_sheet(rows,"202608","202608","unconfirmed.xlsx","test")

    def test_valid_source_does_not_hide_unknown_amount_in_another_upload(self):
        import tempfile
        import hashlib
        import openpyxl
        from entry import execute
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);files=[]
            for index in range(2):
                rows=self.sheet()
                if index:
                    rows[6][0]="未知新增税费";rows[6][2]=10
                book=openpyxl.Workbook()
                for row in rows:book.active.append(row)
                path=root/(str(index)+".xlsx");book.save(path)
                files.append({"local_path":str(path),"name":path.name,"file_id":str(index),"sha256":hashlib.sha256(path.read_bytes()).hexdigest()})
            result=execute({"parameters":{"period":"202602","fetch_kingdee":False},"output_dir":str(root/"output"),"files":{"reports":files}})
        self.assertEqual(result["summary"]["report_count"],1)
        self.assertTrue(any(issue.get("file")=="1.xlsx" for issue in result["issues"]))

    def test_unknown_non_january_cumulative_rejected(self):
        rows=self.sheet();rows[2][2]="本期金额"
        with self.assertRaises(engine.SourceError):self.parse(rows)

    def test_january_month_can_supply_year(self):
        rows=self.sheet("2026年01月");rows[2][2]="本期金额"
        report=self.parse(rows,"202601")
        self.assertEqual(report.values[5,0].amount,report.values[5,1].amount)

    def test_wrong_month_rejected(self):
        with self.assertRaises(engine.SourceError):self.parse(self.sheet("2026年01月"))

    def test_branch_identity_wins_over_head_substring(self):
        self.assertEqual(self.parse(self.sheet()).company,"SHANDONG")

    def test_unknown_nonzero_field_not_discarded(self):
        rows=self.sheet();rows[6][0]="未知新增税费";rows[6][2]=10
        with self.assertRaises(engine.SourceError):self.parse(rows)

    def test_known_unmapped_nonzero_detail_visible(self):
        rows=self.sheet();rows[6][2]=10
        self.assertTrue(any("非零来源明细未映射" in x for x in self.parse(rows).issues))

    def test_same_report_deduplicated_and_conflicts_excluded(self):
        a=self.parse(self.sheet());b=copy.deepcopy(a)
        self.assertEqual(len(engine.select_sources([a,b])[0]),1)
        b.values[5,0].amount+=Decimal(1)
        selected,issues=engine.select_sources([a,b])
        self.assertEqual(selected,[]);self.assertTrue(issues)

    def test_source_precision_and_nonfinite_fail(self):
        for v in ["nan","Infinity","0.001","=1+2",True]:
            with self.subTest(v=v):
                with self.assertRaises(engine.SourceError):engine.amount(v)

    def test_disabled_rule_preserves_management(self):
        target=engine.classify(self.parse(self.sheet()),enabled=False)
        self.assertEqual(target.values[5,1].amount,Decimal(70))
        self.assertEqual(target.rule_version,"none")

    def test_missing_continuing_profit_is_flagged(self):
        self.assertTrue(any("持续经营" in x for x in self.parse(self.sheet()).issues))


    def small_balance(self):
        rows=[["资产负债表"],["甲骨易（湖南）科技有限公司","2026年01月","单位：元"],["资产","行次","期末余额","年初余额","负债","行次","期末余额","年初余额"]]
        for key,label in sorted(engine.CATALOG["small"]["bs"].items(),key=lambda x:int(x[0])):
            row=[""]*8
            offset=0 if int(key)<31 else 4
            row[offset:offset+4]=[label,int(key),0,0]
            if int(key)==41:row[offset]="流动负债合计"
            rows.append(row)
        return rows

    def test_hunan_tech_balance_without_total_colon(self):
        report=engine.parse_sheet(self.small_balance(),"202601","202601","input.xlsx","abc",bs_reclassified=True)
        self.assertEqual(report.company,"HUNAN_TECH")
        self.assertEqual(report.kind,"bs")
        self.assertTrue(all(x["passed"] for x in engine.validate(report)))

    def test_hunan_balance_unknown_total_label_still_rejected(self):
        rows=self.small_balance()
        for row in rows[3:]:
            if row[5]==41:row[4]="未知负债合计"
        with self.assertRaises(engine.SourceError):
            engine.parse_sheet(rows,"202601","202601","input.xlsx","abc")


    def test_detail_numbers_subtotals_and_live_summary_formulas(self):
        import tempfile
        import openpyxl
        from workbook import build_workbook, verify_formulas, TOTAL_RULES
        shandong=self.parse(self.sheet())
        rows=self.sheet()
        rows[1][0]="甲骨易（济南）科技有限公司"
        jinan=self.parse(rows)
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/"test.xlsx"
            audit=build_workbook([shandong,jinan],"202602",path)
            book=openpyxl.load_workbook(path)
            cached=openpyxl.load_workbook(path,data_only=True)
            self.assertEqual(set(book.sheetnames),{
                scope+name+"202602" for scope in ["合并","母公司"]
                for name in ["资产负债表","利润表","现金流量表"]
            })
            expected={}
            company_cells={}
            for item in audit["values"]:
                cell=book[item["sheet"]][item["cell"]]
                self.assertEqual(Decimal(str(cached[item["sheet"]][item["cell"]].value)),Decimal(item["value"]))
                if item["company"] in {"合并","母公司"}:
                    self.assertEqual(cell.data_type,"f")
                    self.assertNotIn("!",cell.value)
                    expected[item["sheet"],item["cell"]]=Decimal(item["value"])
                else:
                    self.assertEqual(cell.data_type,"f" if item["item"] in TOTAL_RULES[item["kind"]] else "n")
                    if cell.data_type=="f":
                        expected[item["sheet"],item["cell"]]=Decimal(item["value"])
                    company_cells[item["sheet"],item["company"],item["item"],item["metric"]]=cell
            self.assertEqual(audit["formula_count"],sum(c.data_type=="f" for sheet in book for row in sheet for c in row))
            self.assertTrue(audit["compilation_notes"])
            self.assertEqual(book.calculation.calcMode,"auto")
            # Editing one company amount must change its group and parent summaries.
            for sheet in ["合并利润表202602","母公司利润表202602"]:
                cell=company_cells[sheet,"SHANDONG",2,0]
                cell.value+=7
                summary=next(v for v in audit["values"] if v["sheet"]==sheet and v["company"] in {"合并","母公司"} and v["item"]==2 and v["metric"]==0)
                expected[sheet,summary["cell"]]+=7
                for item in audit["values"]:
                    if item["sheet"]==sheet and item["metric"]==0 and item["item"] in {19,22,24,42} and item["company"] in {"SHANDONG","合并","母公司"}:
                        expected[sheet,item["cell"]]-=7
            verify_formulas(book,expected)
            book.close();cached.close()

    def test_empty_scope_stays_six_sheets_without_zero_summary(self):
        import tempfile
        import openpyxl
        from workbook import build_workbook
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/"empty.xlsx"
            result=build_workbook([],"202602",path)
            book=openpyxl.load_workbook(path)
            self.assertEqual(len(book.sheetnames),6)
            self.assertEqual(result["formula_count"],0)
            self.assertFalse(result["complete"])
            self.assertTrue(all("不计算为零" in sheet["A3"].value for sheet in book))
            book.close()


    def test_negative_company_and_summary_amounts_use_plain_signed_format(self):
        import tempfile
        import openpyxl
        from workbook import build_workbook
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/"negative.xlsx"
            audit=build_workbook([self.parse(self.sheet())],"202602",path)
            book=openpyxl.load_workbook(path)
            negative_companies=negative_summaries=0
            for value in audit["values"]:
                if Decimal(value["value"])>=0:
                    continue
                cell=book[value["sheet"]][value["cell"]]
                self.assertEqual(cell.number_format.split(";")[1],"-#,##0.00")
                self.assertNotIn("[Red]",cell.number_format)
                self.assertIsNone(cell.font.color)
                if value["company"] in {"合并","母公司"}:
                    negative_summaries+=1
                else:
                    negative_companies+=1
            self.assertGreater(negative_companies,0)
            self.assertGreater(negative_summaries,0)
            book.close()


    def test_balance_totals_exclude_included_details_and_subtract_treasury_stock(self):
        import tempfile
        import openpyxl
        from workbook import build_workbook
        report=engine.parse_sheet(self.small_balance(),"202601","202601","input.xlsx","abc",bs_reclassified=True)
        values={1:160,14:160,34:160,49:100,50:40,51:20,52:20,59:140,60:140,
                61:30,62:10,63:5,64:5,66:20,71:20,72:160}
        for m in [0,1]:
            for i,v in values.items():report.values[i,m].amount=Decimal(v)
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/"test.xlsx"
            audit=build_workbook([report],"202601",path)
            book=openpyxl.load_workbook(path,data_only=True)
            for item in audit["values"]:
                if item["item"] in {59,71,72}:
                    self.assertEqual(book[item["sheet"]][item["cell"]].value,values[item["item"]])
            book.close()

    def test_wrong_source_subtotal_is_not_hidden_by_a_constant_formula(self):
        import tempfile
        from workbook import build_workbook
        report=engine.parse_sheet(self.small_balance(),"202601","202601","input.xlsx","abc",bs_reclassified=True)
        report.values[1,0].amount=Decimal(10)
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError,"公式引用与标准金额不一致"):
                build_workbook([report],"202601",Path(tmp)/"test.xlsx")

    def test_split_files_keep_three_sheets_formulas_values_and_format(self):
        import tempfile
        import openpyxl
        from workbook import build_workbook, split_workbook
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/"甲骨易2026年02月合并报表底稿_部分.xlsx"
            build_workbook([self.parse(self.sheet())],"202602",path)
            full=openpyxl.load_workbook(path)
            cache=openpyxl.load_workbook(path,data_only=True)
            outputs=split_workbook(path)
            self.assertEqual([p.name for p in outputs],["甲骨易2026年02月合并报表_部分.xlsx","甲骨易2026年02月母公司报表_部分.xlsx"])
            for scope,p in zip(["合并","母公司"],outputs):
                book=openpyxl.load_workbook(p)
                values=openpyxl.load_workbook(p,data_only=True)
                self.assertEqual(book.sheetnames,[n for n in full.sheetnames if n.startswith(scope)])
                self.assertEqual(len(book.sheetnames),3)
                self.assertEqual(book.calculation.calcMode,"auto")
                for sheet in book:
                    for row in sheet:
                        for cell in row:
                            original=full[sheet.title][cell.coordinate]
                            self.assertEqual(cell.value,original.value)
                            self.assertEqual(cell._style,original._style)
                            self.assertEqual(values[sheet.title][cell.coordinate].value,cache[sheet.title][cell.coordinate].value)
                    self.assertEqual(sheet.freeze_panes,full[sheet.title].freeze_panes)
                    self.assertEqual(str(sheet.print_area),str(full[sheet.title].print_area))
                book.close();values.close()
            full.close();cache.close()

if __name__=="__main__":unittest.main()
