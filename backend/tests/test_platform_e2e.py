from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

from fastapi.testclient import TestClient
from helpers import auth_client
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app import model_service, orchestrator
from app.database import SessionLocal
from app.models import ModelConnection
from app.registry import registry
from app.worker import run_once


def workbook_bytes(headers: list[str], rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


class FakeVerificationResponse:
    """符合 Tool Calling 验证要求的成功响应结构。"""

    def raise_for_status(self) -> None:
        return None

    def json(self) -> dict[str, object]:
        return {
            "choices": [
                {
                    "message": {
                        "role": "assistant",
                        "content": "",
                        "tool_calls": [
                            {
                                "id": "call_1",
                                "type": "function",
                                "function": {
                                    "name": "tool_call_supported",
                                    "arguments": "{}",
                                },
                            }
                        ],
                    }
                }
            ]
        }


def upload(client: TestClient, role: str, name: str, content: bytes) -> str:
    response = client.post(
        "/api/files",
        data={"role": role},
        files={
            "upload": (
                name,
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200, response.text
    return response.json()["id"]


def test_registry_and_admin_boundary() -> None:
    with auth_client() as client:
        skills = client.get("/api/skills")
        assert skills.status_code == 200
        assert [item["id"] for item in skills.json()] == [
            "labor-invoice-check",
            "compliance-spot-check",
            "reconcile-bank",
            "split-by-sales",
            "receivables-merge",
            "dreame-ar-progress-diff",
            "withholding-report-rename",
            "order-daily-summary",
            "project-detail-to-ledger",
            "dept-expense-alloc",
        ]

        denied = client.post("/api/admin/registry/reload")
        assert denied.status_code == 403

    with auth_client(role="skill_admin") as admin_client:
        allowed = admin_client.post("/api/admin/registry/reload")
        assert allowed.status_code == 200
        assert allowed.json()["skills"] == 19
        assert allowed.json()["errors"] == []


def test_qwen_tool_call_disables_thinking(monkeypatch) -> None:
    registry.refresh()
    skill = registry.get("reconcile-bank")
    assert skill is not None
    captured: dict[str, object] = {}

    class FakeResponse:
        def raise_for_status(self) -> None:
            return None

        def json(self) -> dict[str, object]:
            return {
                "usage": {"prompt_tokens": 12, "completion_tokens": 5},
                "choices": [
                    {
                        "message": {
                            "tool_calls": [
                                {
                                    "function": {
                                        "arguments": (
                                            '{"amount_tolerance":3,"date_tolerance_days":4}'
                                        )
                                    }
                                }
                            ]
                        }
                    }
                ]
            }

    def fake_post(*_, **kwargs):
        captured.update(kwargs["json"])
        return FakeResponse()

    monkeypatch.setattr(
        orchestrator,
        "settings",
        SimpleNamespace(
            llm_base_url="https://dashscope.aliyuncs.com/compatible-mode/v1",
            llm_api_key="test-key",
            llm_model="qwen3.7-plus",
        ),
    )
    monkeypatch.setattr(orchestrator.httpx, "post", fake_post)

    trace: dict[str, int | str] = {}
    parameters, missing, source, notes = orchestrator.interpret_parameters(
        skill,
        "金额差异 3 元以内，日期相差 4 天可以匹配",
        {},
        trace=trace,
    )
    assert captured["enable_thinking"] is False
    assert parameters == {"amount_tolerance": 3, "date_tolerance_days": 4}
    assert missing == []
    assert source == "llm"
    assert notes == []
    assert trace["status"] == "succeeded"
    assert trace["input_tokens"] == 12
    assert trace["output_tokens"] == 5
    assert int(trace["duration_ms"]) >= 0


def test_api_key_auto_detection_and_model_selection(monkeypatch) -> None:
    api_key = "sk-test-model-connection-secret"

    class FakeModelsResponse:
        def raise_for_status(self) -> None:
            return None

        def json(self) -> dict[str, object]:
            return {
                "data": [
                    {"id": "qwen3.7-plus"},
                    {"id": "qwen3.6-plus"},
                    {"id": "qwen-image-2.0-pro"},
                ]
            }

    monkeypatch.setattr(
        model_service.httpx,
        "get",
        lambda *_, **__: FakeModelsResponse(),
    )
    monkeypatch.setattr(
        model_service.httpx,
        "post",
        lambda *_, **__: FakeVerificationResponse(),
    )

    with auth_client() as client:
        connected = client.post("/api/model-connections", json={"api_key": api_key})
        assert connected.status_code == 200, connected.text
        body = connected.json()
        connection_id = body["id"]
        assert body["provider"] == "qwen"
        assert body["selected_model"] == "qwen3.7-plus"
        assert body["models"] == ["qwen3.7-plus", "qwen3.6-plus"]
        assert api_key not in connected.text
        assert body["api_key_hint"].endswith("cret")

        listed = client.get("/api/model-connections")
        assert listed.status_code == 200
        assert [item["id"] for item in listed.json()] == [connection_id]

        selected = client.patch(
            f"/api/model-connections/{connection_id}",
            json={"selected_model": "qwen3.6-plus"},
        )
        assert selected.status_code == 200
        assert selected.json()["selected_model"] == "qwen3.6-plus"

        with SessionLocal() as db:
            stored = db.scalar(select(ModelConnection).where(ModelConnection.id == connection_id))
            assert stored is not None
            assert api_key not in stored.api_key_encrypted

        removed = client.delete(f"/api/model-connections/{connection_id}")
        assert removed.status_code == 204
        assert client.get("/api/model-connections").json() == []


def test_upload_run_worker_and_download(monkeypatch) -> None:
    bank = workbook_bytes(
        ["交易日期", "交易金额", "流水号"],
        [
            ["2026-07-01", 100.00, "B-001"],
            ["2026-07-02", 200.00, "B-002"],
            ["2026-07-10", 999.00, "B-003"],
        ],
    )
    ledger = workbook_bytes(
        ["凭证日期", "金额", "凭证号"],
        [
            ["2026-07-01", 100.00, "记-0001"],
            ["2026-07-03", 200.50, "记-0002"],
            ["2026-07-20", 700.00, "记-0003"],
        ],
    )

    class FakeModelsResponse:
        def raise_for_status(self) -> None:
            return None

        def json(self) -> dict[str, object]:
            return {"data": [{"id": "qwen3.7-plus"}]}

    class FakeToolResponse:
        def raise_for_status(self) -> None:
            return None

        def json(self) -> dict[str, object]:
            return {
                "choices": [
                    {
                        "message": {
                            "tool_calls": [
                                {
                                    "function": {
                                        "arguments": (
                                            '{"amount_tolerance":1,"date_tolerance_days":2}'
                                        )
                                    }
                                }
                            ]
                        }
                    }
                ]
            }

    def fake_post(url, **kwargs):
        del url
        payload = kwargs.get("json") or {}
        tools = payload.get("tools") or []
        if tools and tools[0].get("function", {}).get("name") == "tool_call_supported":
            return FakeVerificationResponse()
        return FakeToolResponse()

    monkeypatch.setattr(
        model_service.httpx,
        "get",
        lambda *_, **__: FakeModelsResponse(),
    )
    monkeypatch.setattr(
        orchestrator.httpx,
        "post",
        fake_post,
    )

    with auth_client() as client:
        connection = client.post(
            "/api/model-connections",
            json={"api_key": "sk-e2e-model-secret"},
        )
        assert connection.status_code == 200
        connection_id = connection.json()["id"]
        bank_id = upload(client, "bank_file", "银行流水.xlsx", bank)
        ledger_id = upload(client, "ledger_file", "财务总账.xlsx", ledger)

        created = client.post(
            "/api/runs",
            json={
                "skill_id": "reconcile-bank",
                "message": "金额差异 1 元以内、日期相差 2 天可以匹配",
                "parameters": {},
                "files": {
                    "bank_file": bank_id,
                    "ledger_file": ledger_id,
                },
                "idempotency_key": "e2e-reconcile-001",
                "model_connection_id": connection_id,
                "model": "qwen3.7-plus",
            },
        )
        assert created.status_code == 200, created.text
        run_id = created.json()["id"]
        assert created.json()["state"] == "waiting_confirmation"
        assert created.json()["model_provider"] == "qwen"
        assert created.json()["model_name"] == "qwen3.7-plus"
        assert created.json()["parameters"] == {
            "amount_tolerance": 1.0,
            "date_tolerance_days": 2,
        }

        confirmed = client.post(f"/api/runs/{run_id}/confirm")
        assert confirmed.status_code == 200, confirmed.text
        assert confirmed.json()["state"] == "queued"

        assert run_once(("python",)) is True

        completed = client.get(f"/api/runs/{run_id}")
        assert completed.status_code == 200
        result = completed.json()
        assert result["state"] == "succeeded", result["error_message"]
        assert result["progress"] == 100
        assert result["result"]["summary"] == {
            "bank_records": 3,
            "ledger_records": 3,
            "matched": 2,
            "unmatched_bank": 1,
            "unmatched_ledger": 1,
        }

        artifact = result["result"]["output_files"][0]
        downloaded = client.get(artifact["download_url"])
        assert downloaded.status_code == 200
        workbook = load_workbook(BytesIO(downloaded.content), data_only=True)
        assert workbook.sheetnames == ["匹配明细", "银行未匹配", "总账未匹配"]
        assert workbook["匹配明细"].max_row == 3
        assert workbook["银行未匹配"].max_row == 2
        assert workbook["总账未匹配"].max_row == 2
        workbook.close()

        events = client.get(
            f"/api/runs/{run_id}/events",
            params={"user_id": "demo-user", "department_id": "finance"},
        )
        assert events.status_code == 200
        assert "任务执行完成" in events.text
