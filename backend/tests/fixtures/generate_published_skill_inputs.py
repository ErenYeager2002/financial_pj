from __future__ import annotations

import argparse
import csv
import json
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill


def save_rows(path: Path, sheet_name: str, rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def compliance_inputs(root: Path) -> dict[str, list[str]]:
    directory = root / "compliance-spot-check"
    directory.mkdir()
    headers = [
        "年度",
        "销售人员",
        "客户名称",
        "新智云单号",
        "文件名",
        "应收金额",
        "交付月份",
        "账龄(月份）",
        "结算阶段",
        "0604销售预计回款日期",
        "销售解释说明",
        "有无合同",
        "合同分类",
        "框架合同是否存在PO单",
        "应收金额是否有客户正式确认",
        "客户结算周期",
        "是否按月给客户发结算单",
    ]
    data = [
        [2026, "测试销售甲", "测试客户甲", "TEST-SO-001", "测试项目1", 280000, "202412", 8],
        [2026, "测试销售甲", "测试客户乙", "TEST-SO-002", "测试项目2", 50000, "202501", 7],
        [2026, "测试销售乙", "测试客户丙", "TEST-SO-003", "测试项目3", 3200, "202411", 9],
        [2026, "测试销售乙", "测试客户丁", "TEST-SO-004", "测试项目4", 1500, "202412", 8],
        [2026, "测试销售丙", "测试客户戊", "TEST-SO-005", "测试项目5", 18000, "202410", 10],
        [2026, "测试销售丙", "测试客户己", "TEST-SO-006", "测试项目6", 12000, "202411", 9],
    ]
    rows = [headers] + [row + [""] * 9 for row in data]
    receivables = directory / "合成应收all.xlsx"
    save_rows(receivables, "2026.8.27", rows)
    history = directory / "合成抽查历史.csv"
    with history.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["销售", "客户", "交付月份", "抽查日期", "反馈状态"])
        writer.writerow(["测试销售丙", "测试客户戊", "202410", "2026-08-20", "未反馈"])
    return {"receivables": [str(receivables)], "history": [str(history)]}


def department_inputs(root: Path) -> dict[str, list[str]]:
    directory = root / "dept-expense-alloc"
    directory.mkdir()

    balance = directory / "01_文化主体余额表.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "发生额及余额表"
    sheet["A2"] = "文化"
    sheet["A3"] = "科目编码"
    sheet["B3"] = "科目名称"
    sheet["C3"] = "期初余额"
    sheet["E3"] = "本期发生"
    sheet["G3"] = "期末余额"
    sheet["C4"] = "借方"
    sheet["D4"] = "贷方"
    sheet["E4"] = "借方"
    sheet["F4"] = "贷方"
    for row, code, name, debit, credit in [
        (5, "5101", "主营业务收入", None, 1000),
        (6, "540103", "翻译语言服务", 200, None),
        (7, "540109", "工资", 100, None),
        (8, "5402", "税金及附加", 50, None),
        (9, "5502", "管理费用", 80, None),
        (10, "5504", "财务费用", 10, None),
    ]:
        sheet.cell(row, 1, code)
        sheet.cell(row, 2, name)
        sheet.cell(row, 5, debit)
        sheet.cell(row, 6, credit)
    workbook.save(balance)
    workbook.close()

    people = directory / "04_人员归属表.xlsx"
    save_rows(
        people,
        "人员归属",
        [
            ["姓名", "组织架构1", "组织架构2"],
            ["测试人员甲", "营销中心", "本地化"],
            ["测试人员乙", "运营保障中心", "运营保障中心"],
            ["测试人员丙", "项目中心", "项目一组"],
        ],
    )
    income = directory / "03_收入底稿.xlsx"
    save_rows(income, "收入", [["申请人", "金额"], ["测试人员甲", 600], ["测试人员甲", 400]])
    expenses = directory / "05_用友按人明细.xlsx"
    save_rows(
        expenses,
        "按人明细",
        [
            ["姓名", "科目编码", "科目名称", "金额"],
            ["测试人员丙", "540109", "工资", 100],
            ["测试人员乙", "5502", "管理费用", 80],
        ],
    )
    return {"materials": [str(balance), str(income), str(people), str(expenses)]}


def dreame_inputs(root: Path) -> dict[str, list[str]]:
    directory = root / "dreame-ar-progress-diff"
    directory.mkdir()

    def write_version(path: Path, amount: float, color: str, note: str) -> None:
        workbook = Workbook()
        workbook.active.title = "其他"
        sheet = workbook.create_sheet("应收进度")
        sheet.cell(1, 1, "合成测试数据")
        sheet.cell(2, 2, "测试业务线（测试人员）")
        sheet.cell(2, 3, "PO时间")
        sheet.cell(2, 4, "发票上传时间")
        sheet.cell(2, 5, "预计付款时间")
        sheet.cell(3, 1, "7月")
        sheet.cell(3, 2, amount).fill = PatternFill("solid", fgColor=color)
        sheet.cell(3, 3, note)
        sheet.cell(4, 1, "为方便统计金额：绿色待付款")
        workbook.save(path)
        workbook.close()

    older = directory / "合成进度-0820.xlsx"
    newer = directory / "合成进度-0827.xlsx"
    write_version(older, 1000, "FFFF00", "0820发送对账单")
    write_version(newer, 1000, "FED4A4", "0827已催促")
    return {"versions": [str(older), str(newer)]}


def labor_inputs(root: Path) -> dict[str, list[str]]:
    directory = root / "labor-invoice-check"
    directory.mkdir()
    labor = directory / "合成劳务清单.xlsx"
    workbook = Workbook()
    workbook.active.title = "说明"
    workbook.active["A1"] = "合成测试数据"
    sheet = workbook.create_sheet("个人译费2026")
    sheet.append(["2026年8月国内个人译费", None, None, None, None, None])
    sheet.append(["序号", "供应商姓名", "应付金额", "备注", "开户名", "身份证号/护照号"])
    sheet.append([1, "测试人员甲", 1500, "Freelancer", "测试人员甲", "TEST-ID-001"])
    sheet.append([2, "测试实习生", 500, "Intern", "测试实习生", "TEST-ID-002"])
    workbook.save(labor)
    workbook.close()

    invoice = directory / "合成发票台账.xlsx"
    save_rows(
        invoice,
        "发票明细202608",
        [
            ["序号", "销售方信息名称", "销售方信息纳税人识别号", "金额（元）", "合计金额（元）"],
            [1, "测试人员甲", "TEST-ID-001", 1500, 1500],
        ],
    )
    return {"labor_list": [str(labor)], "invoice_ledger": [str(invoice)]}


def order_inputs(root: Path) -> dict[str, list[str]]:
    directory = root / "order-daily-summary"
    directory.mkdir()
    orders = directory / "合成下单明细.xlsx"
    save_rows(
        orders,
        "下单",
        [
            ["销售", "SO", "订单名称", "下单日期", "下单预估额/本币"],
            ["测试销售甲", "TEST-SO-101", "测试项目甲", "2026-08-26", 10000],
            ["测试销售乙", "TEST-SO-102", "测试项目乙", "2026-08-27", 5000],
        ],
    )
    return {"orders": [str(orders)]}


def project_inputs(root: Path) -> dict[str, list[str]]:
    directory = root / "project-detail-to-ledger"
    directory.mkdir()
    project = directory / "合成项目明细.xlsx"
    save_rows(
        project,
        "项目明细",
        [
            [
                "销售",
                "客户",
                "SO",
                "SOD",
                "业务类别",
                "订单名称",
                "下单日期",
                "整单交付日期",
                "下单数量",
                "单价",
                "交付额/本币",
                "项目经理",
            ],
            [
                "测试销售甲",
                "测试客户甲",
                "TEST-SO-NEW",
                "TEST-SOD-NEW",
                "笔译",
                "测试项目甲",
                "2026-08-26",
                "2026-08-27",
                100,
                2.5,
                250,
                "测试经理甲",
            ],
        ],
    )
    ledger = directory / "合成盈亏核算表.xlsx"
    save_rows(
        ledger,
        "明细",
        [
            [
                "销售人员",
                "客户名称",
                "新智云单号",
                "翻译类型",
                "文件名",
                "项目下单日期",
                "项目交付日期",
                "字数统计",
                "价格",
                "应收金额",
                "实收金额",
                "项目经理",
            ],
            [
                "测试销售乙",
                "测试客户旧",
                "TEST-SO-OLD",
                "笔译",
                "测试旧项目",
                "2026-08-01",
                "2026-08-02",
                50,
                2,
                100,
                "TEST-SOD-OLD",
                "测试经理乙",
            ],
        ],
    )
    return {"project_detail": [str(project)], "ledger": [str(ledger)]}


def receivables_inputs(root: Path) -> dict[str, list[str]]:
    directory = root / "receivables-merge-and-split"
    directory.mkdir()
    source = directory / "合成本期应收源台账.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    for year in ("2026", "2025"):
        sheet = workbook.create_sheet(year)
        sheet.append(
            ["销售人员", "客户名称", "单号", "新智云单号", "文件名", "应收金额", "项目交付"]
        )
        sheet.append(
            [
                "测试销售甲",
                "测试客户甲",
                f"TEST-{year}-1",
                f"TEST-{year}-1",
                "测试项目",
                1000,
                f"{year}03",
            ]
        )
        sheet.append(
            [
                "测试销售乙",
                "测试客户乙",
                f"TEST-{year}-2",
                f"TEST-{year}-2",
                "测试项目",
                0,
                f"{year}04",
            ]
        )
    batch = workbook.create_sheet("8月批量")
    batch.append(["销售", "客户", "订单号", "名称", "完成时间", "订单折合本币"])
    batch.append(["测试销售丙", "测试客户丙", "TEST-GM-1", "测试项目", "2026-08-20", 500])
    workbook.save(source)
    workbook.close()
    return {"primary": [str(source)]}


def reconciliation_inputs(root: Path) -> dict[str, list[str]]:
    directory = root / "reconcile-bank"
    directory.mkdir()
    bank = directory / "合成银行流水.xlsx"
    ledger = directory / "合成财务总账.xlsx"
    save_rows(
        bank,
        "银行流水",
        [
            ["交易日期", "交易金额", "流水号"],
            ["2026-08-20", 100, "TEST-B-001"],
            ["2026-08-21", 200, "TEST-B-002"],
            ["2026-08-22", 999, "TEST-B-003"],
        ],
    )
    save_rows(
        ledger,
        "财务总账",
        [
            ["凭证日期", "金额", "凭证号"],
            ["2026-08-20", 100, "TEST-L-001"],
            ["2026-08-22", 200.5, "TEST-L-002"],
            ["2026-08-25", 700, "TEST-L-003"],
        ],
    )
    return {"bank_file": [str(bank)], "ledger_file": [str(ledger)]}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    root = (args.output or Path(tempfile.mkdtemp(prefix="financial_skill_e2e_"))).resolve()
    root.mkdir(parents=True, exist_ok=True)
    manifest = {
        "compliance-spot-check": compliance_inputs(root),
        "dept-expense-alloc": department_inputs(root),
        "dreame-ar-progress-diff": dreame_inputs(root),
        "labor-invoice-check": labor_inputs(root),
        "order-daily-summary": order_inputs(root),
        "project-detail-to-ledger": project_inputs(root),
        "receivables-merge-and-split": receivables_inputs(root),
        "reconcile-bank": reconciliation_inputs(root),
    }
    manifest_path = root / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(str(manifest_path))


if __name__ == "__main__":
    main()
