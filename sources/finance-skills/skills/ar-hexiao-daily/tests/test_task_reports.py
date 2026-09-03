# -*- coding: utf-8 -*-
import json
import subprocess
import sys
from pathlib import Path

import openpyxl

import build_task_reports as B
import workbook_finalize as W


def _find_lightweight_checker():
    for parent in Path(__file__).resolve().parents:
        for relative in (
            ("tools", "xlsx_lightweight_audit.py"),
            ("03_tools", "scripts", "xlsx_lightweight_audit.py"),
        ):
            candidate = parent.joinpath(*relative)
            if candidate.is_file():
                return candidate
    return None


def _report(path, title, value):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    ws.append(["项目", "值"])
    ws.append(["日期", value])
    wb.save(path)


def _daily_report(path, value):
    workbook = openpyxl.Workbook()
    for index, title in enumerate(("今日清单", "跨月计提补填", "流转表怎么填")):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = title
        sheet.append(["项目", "值"])
        sheet.append([title, value])
    workbook.save(path)


def test_build_task_reports_keeps_one_integrated_static_range_file(tmp_path):
    out = tmp_path / "04_产出"
    out.mkdir()
    for token, date in (("20260808", "2026-08-08"), ("20260809", "2026-08-09")):
        _daily_report(out / f"核销日清_{token}.xlsx", date)
        for prefix in ("变更清单", "订单写入差异"):
            _report(out / f"{prefix}_{token}.xlsx", "内部", date)
        (out / f"判定结果_{token}.json").write_text(
            json.dumps({
                "payment_count": 2,
                "counts": {"total": 3, "auto": 2, "hold": 1, "exception": 0},
            }, ensure_ascii=False),
            encoding="utf-8",
        )
    _report(out / "变更清单_20260808_2025.xlsx", "内部", "2026-08-08")
    _report(out / "核销日清_20260810.xlsx", "范围外", "2026-08-10")

    outputs = B.build(tmp_path, "2026-08-08", "2026-08-09")

    assert len(outputs) == 1
    path = outputs[0]
    assert path.name == "核销日清_20260808_20260809.xlsx"
    wb = openpyxl.load_workbook(path, data_only=False)
    assert wb.sheetnames == ["任务范围", "核销明细", "跨月计提补填", "流转表怎么填"]
    assert {row[0].value for row in wb["任务范围"].iter_rows(min_row=2)} == {
        "2026-08-08", "2026-08-09"
    }
    detail = wb["核销明细"]
    assert detail.cell(row=1, column=1).value == "核销日期"
    assert [detail.cell(row=row, column=1).value for row in range(2, 4)] == [
        "2026-08-08", "2026-08-09"
    ]
    for title in ("跨月计提补填", "流转表怎么填"):
        assert [wb[title].cell(row=row, column=1).value for row in range(2, 4)] == [
            "2026-08-08", "2026-08-09"
        ]
    wb.close()
    audit = W.inspect_calculation(path)
    assert audit.formula_cells == 0
    assert audit.full_calc_on_load == "0"
    assert audit.force_full_calc == "0"
    checker = _find_lightweight_checker()
    if checker is not None:
        checked = subprocess.run(
            [sys.executable, str(checker), str(path), "--strict"],
            capture_output=True, text=True, encoding="utf-8",
        )
        assert checked.returncode == 0, checked.stdout + checked.stderr

    assert not list(out.glob("核销日清_2026080[89].xlsx"))
    assert not list(out.glob("变更清单_2026080[89].xlsx"))
    assert not list(out.glob("订单写入差异_2026080[89].xlsx"))
    assert not (out / "变更清单_20260808_2025.xlsx").exists()
    assert (out / "核销日清_20260810.xlsx").is_file()
    assert (out / "判定结果_20260808.json").is_file()
    assert (out / "判定结果_20260809.json").is_file()


def test_build_task_reports_keeps_formula_like_text_static(tmp_path):
    out = tmp_path / "04_产出"
    out.mkdir()
    for token in ("20260808", "20260809"):
        _report(out / f"核销日清_{token}.xlsx", "今日清单", "=说明文字")
        (out / f"判定结果_{token}.json").write_text(
            json.dumps({"payment_count": 0, "counts": {}}),
            encoding="utf-8",
        )

    output = B.build(tmp_path, "2026-08-08", "2026-08-09")[0]

    workbook = openpyxl.load_workbook(output, data_only=False)
    assert workbook["核销明细"].cell(row=2, column=3).value == "=说明文字"
    assert workbook["核销明细"].cell(row=2, column=3).data_type == "s"
    assert W.inspect_calculation(output).formula_cells == 0
    workbook.close()


def test_build_task_reports_accepts_confirmed_empty_fetch_day(tmp_path):
    out = tmp_path / "04_产出"
    out.mkdir()
    _report(out / "核销日清_20260809.xlsx", "今日清单", "2026-08-09")
    (out / "判定结果_20260809.json").write_text(
        json.dumps(
            {"payment_count": 1, "counts": {"total": 1, "auto": 1, "hold": 0, "exception": 0}},
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    outputs = B.build(
        tmp_path,
        "2026-08-08",
        "2026-08-09",
        empty_dates=["2026-08-08"],
    )

    assert len(outputs) == 1
    workbook = openpyxl.load_workbook(outputs[0], data_only=True)
    rows = list(workbook["任务范围"].iter_rows(min_row=2, values_only=True))
    assert rows[0][:2] == ("2026-08-08", "无核销记录，已跳过")
    assert rows[1][:2] == ("2026-08-09", "已纳入")
    workbook.close()


def test_build_task_reports_uses_only_explicit_sparse_dates(tmp_path):
    out = tmp_path / "04_产出"
    out.mkdir()
    for token, date in (("20260808", "2026-08-08"), ("20260810", "2026-08-10")):
        _report(out / f"核销日清_{token}.xlsx", "今日清单", date)
        (out / f"判定结果_{token}.json").write_text(
            json.dumps({"payment_count": 1, "counts": {"total": 1}}, ensure_ascii=False),
            encoding="utf-8",
        )

    outputs = B.build(
        tmp_path,
        "2026-08-08",
        "2026-08-10",
        selected_dates=["2026-08-08", "2026-08-10"],
    )

    assert outputs[0].name == "核销日清_已选2日_20260808_20260810.xlsx"
    workbook = openpyxl.load_workbook(outputs[0], data_only=True)
    assert [row[0] for row in workbook["任务范围"].iter_rows(min_row=2, values_only=True)] == [
        "2026-08-08",
        "2026-08-10",
    ]
    workbook.close()


def test_build_task_reports_single_day_keeps_only_integrated_workbook(tmp_path):
    out = tmp_path / "04_产出"
    out.mkdir()
    _report(out / "核销日清_20260808.xlsx", "今日清单", "2026-08-08")
    _report(out / "变更清单_20260808.xlsx", "内部", "2026-08-08")
    _report(out / "订单写入差异_20260808.xlsx", "内部", "2026-08-08")

    outputs = B.build(
        tmp_path,
        "2026-08-08",
        "2026-08-08",
        selected_dates=["2026-08-08"],
    )

    assert [path.name for path in outputs] == ["核销日清_20260808_20260808.xlsx"]
    assert sorted(path.name for path in out.glob("*.xlsx")) == [
        "核销日清_20260808_20260808.xlsx"
    ]
