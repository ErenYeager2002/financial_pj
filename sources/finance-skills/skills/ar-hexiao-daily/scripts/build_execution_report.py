#!/usr/bin/env python3
"""Join initial decisions, actual writes and post-write outcomes without reapplying."""
from __future__ import annotations

import argparse
import collections
import json
import shutil
from pathlib import Path

import common
import workbook_finalize
from build_execution_evidence import digest, owned_file, record_identity
from execution_lineage import checked_records, indexed_decisions, match_final_records


def final_outcome(after: dict, checks: dict) -> dict:
    key = after["case_id"]
    checked = checks["conflict"].get(key) or checks["skip"].get(key) or after
    check = checked.get("_check") or {}
    reason = check.get("reason") or after.get("reason") or ""
    if not str(reason).strip():
        raise ValueError("最终原因缺失，不能发布只有状态的记录")
    return {
        "case_id": key, "sod": after.get("sod"), "bucket": after["bucket"],
        "status": "conflict" if key in checks["conflict"] else "skipped" if after["bucket"] == "auto" else after["bucket"],
        "code": check.get("code") or after.get("code") or "", "reason": reason,
        "reason_detail": {
            "stage": "post_write_validation", "classification_code": after.get("code") or "",
            "check": check,
            "facts": {field: after[field] for field in (
                "current_values", "idempotence_audit", "sod_capacity_audit", "ambiguous_sod_waterfall",
                "split_payment_source", "warning_codes", "delivery_date_issue", "ledger_year",
                "ledger_row_ref", "same_so_multi_sod_absorbed", "tail_tolerance_absorbed",
                "split_chain_group_id", "split_chain_index", "split_chain_count",
            ) if field in after},
            "conclusion": reason,
        },
    }


def build(workspace: Path, checked: Path) -> dict:
    plan = json.loads(checked.read_text(encoding="utf-8"))
    day = common.norm_date(plan.get("hexiao_date"))
    if day is None:
        raise ValueError("最终报告缺少核销日期")
    tag = day.strftime("%Y%m%d")
    output = workspace / "04_产出"
    initial_payload = json.loads((output / f"判定结果_{tag}.json").read_text(encoding="utf-8"))
    initial = indexed_decisions(initial_payload)
    review_dir = workspace / "execution-review" / "04_产出"
    final_payload = json.loads((review_dir / f"判定结果_{tag}.json").read_text(encoding="utf-8"))
    final = indexed_decisions(final_payload)
    rechecked = json.loads((review_dir / f"写入计划_校验后_{tag}.json").read_text(encoding="utf-8"))
    if any(common.norm_date(item.get("hexiao_date")) != day for item in (initial_payload, final_payload, rechecked)):
        raise ValueError("首次判定、写后判定或校验计划与固定核销日期不一致")
    if rechecked.get("write"):
        raise ValueError("写后仍有待写项目，不能生成完成报告")
    initial_checks = checked_records(plan, {key for key, row in initial.items() if row["bucket"] == "auto"})
    final_checks = checked_records(rechecked, {key for key, row in final.items() if row["bucket"] == "auto"})
    links = match_final_records(initial, final)
    verification = json.loads((output / "写后业务核对.json").read_text(encoding="utf-8"))
    if not verification.get("verified"):
        raise ValueError("盈亏实际执行核对没有通过")
    written = set(verification.get("verified_case_ids") or [])
    if written != {str(row.get("case_id")) for row in plan.get("write") or []}:
        raise ValueError("实际执行项目与首次校验后计划不一致")
    initial_skip = initial_checks["skip"]
    conflicts = initial_checks["conflict"]
    if any(key not in conflicts for key, targets in links.items()
           if any(target in final_checks["conflict"] for target in targets)):
        raise ValueError("写后出现首次计划之外的冲突，不能发布")
    outcomes = {key: final_outcome(row, final_checks) for key, row in final.items()}
    parents = collections.defaultdict(list)
    for key, targets in links.items():
        for target in targets:
            parents[target].append(key)
    records = []
    for key, before in initial.items():
        components = [outcomes[target] for target in links[key]]
        # Counts remain in initial-record units. Mixed descendants keep every
        # outcome below and use the most restrictive disposition in the summary.
        status = next(value for value in ("exception", "conflict", "hold", "skipped")
                      if any(item["status"] == value for item in components))
        if status == "skipped" and key in written:
            status = "completed"
        execute_status = "written_verified" if key in written else "skipped" if key in initial_skip else "conflict" if key in conflicts else "not_executed"
        reason = components[0]["reason"] if len(components) == 1 else "；".join(
            f"{item['case_id']}：{item['reason']}" for item in components)
        changed = any(before["bucket"] != item["bucket"] or (before.get("code") or "") != item["code"]
                      for item in components)
        mapping_changed = links[key] != [key] or len(parents[links[key][0]]) > 1
        execution_label = {"written_verified": "本条已写入并回读", "skipped": "本条首次校验跳过，未独立写入",
                           "conflict": "本条首次校验冲突，未写入", "not_executed": "本条未独立执行写入"}[execute_status]
        initial_check = (initial_checks["write"].get(key) or initial_skip.get(key) or conflicts.get(key) or {}).get("_check") or {}
        execution_reason = initial_check.get("reason") or before.get("reason") or ""
        if not str(execution_reason).strip():
            raise ValueError("首次执行或跳过原因缺失，不能仅报告写后状态")
        change_reason = (f"首次 {before.get('code') or before['bucket']}：{before.get('reason') or ''}；"
                         f"写后复核：{reason}。{execution_label}。") if changed or mapping_changed else ""
        records.append({
            "record_id": record_identity(before), "case_id": key,
            "ar": before.get("ar"), "so": before.get("so"), "sod": before.get("sod"),
            "initial_bucket": before["bucket"], "initial_code": before.get("code") or "",
            "initial_reason": before.get("reason") or "",
            "execution_status": execute_status, "final_status": status,
            "execution_reason": execution_reason, "execution_check": initial_check,
            "execution_evidence": (verification.get("execution_rows") or {}).get(key) or {},
            "final_code": components[0]["code"] if len(components) == 1 else "MULTIPLE_FINAL_OUTCOMES", "final_reason": reason,
            "final_outcomes": components,
            "lineage": {"source": before["source_lineage"], "initial_case_id": key,
                        "final_case_ids": links[key], "mapping_changed": mapping_changed,
                        "shared_initial_case_ids": sorted({parent for target in links[key] for parent in parents[target]})},
            "reason_detail": {
                "stage": "post_write_validation",
                "outcomes": [item["reason_detail"] for item in components],
                "initial_evidence_id": record_identity(before),
                "conclusion": reason,
            },
            "state_changed": changed, "change_reason": change_reason,
        })
    counts = dict(collections.Counter(row["final_status"] for row in records))
    counts["total"] = len(records)
    counts["written_records"] = len(written)
    counts["written_orders"] = len({str(initial[key].get("so")) for key in written if initial[key].get("so")})
    counts["initial_skipped_records"] = len(initial_skip)
    counts["final_classified_records"] = len(final)
    flow = json.loads((output / "流转阶段执行结果.json").read_text(encoding="utf-8"))
    holds = json.loads((output / "挂账重扫结果.json").read_text(encoding="utf-8"))
    if common.norm_date(holds.get("reconciliation_date")) != day:
        raise ValueError("挂账重扫结果与最终报告日期不一致")
    baseline = holds.get("baseline") or {}
    if baseline.get("existed"):
        hold_source = owned_file(workspace, workspace / baseline["snapshot"])
        if digest(hold_source) != baseline.get("snapshot_sha256"):
            raise ValueError("最终报告中的历史挂账来源与重扫前快照不一致")
    for name, fingerprint in (holds.get("sources") or {}).items():
        if digest(owned_file(workspace, workspace / name)) != fingerprint:
            raise ValueError("挂账重扫来源发生变化，不能与最终报告混用")
    hold_by_record = collections.defaultdict(list)
    known_record_ids = {row["record_id"] for row in records}
    for hold in holds.get("records") or []:
        for identity in hold.get("initial_record_ids") or []:
            if identity not in known_record_ids:
                raise ValueError("挂账重扫引用了本任务首次判定之外的记录")
            hold_by_record[identity].append(hold)
    for row in records:
        row["hold_rescan"] = hold_by_record[row["record_id"]]
    return {"schema_version": "ar-final-result-v1", "reconciliation_date": day.isoformat(),
            "counts": counts, "records": records, "write_count": len(written),
            "count_basis": "initial_decision_records", "final_classified_counts": dict(collections.Counter(row["bucket"] for row in final.values())),
            "skip_count": len(initial_skip), "flow": flow, "hold_rescan": holds,
            "post_write_validation": rechecked.get("counts") or {},
            "publication_note": "此报告记录已复核的业务结果；正式材料版本与发布状态以平台执行记录为准。"}


def write_report(payload: dict, target: Path, initial_result: dict, checked_plan: dict) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from build_worklist import HEADERS, _row

    initial_report = target.with_name(f"首次{target.name}")
    if not initial_report.exists():
        if not target.is_file():
            raise ValueError("首次日清不存在，不能替代原有业务明细")
        shutil.copy2(target, initial_report)
    workbook = openpyxl.load_workbook(initial_report)
    if "流转表怎么填" not in workbook.sheetnames:
        workbook.close()
        raise ValueError("首次日清缺少流转填写说明")
    for sheet in list(workbook):
        if sheet.title != "流转表怎么填":
            workbook.remove(sheet)
    summary = workbook.create_sheet("任务范围", 0)
    summary.append(["项目", "结果"])
    summary.append(["核销日期", payload["reconciliation_date"]])
    summary.append(["报告阶段", "写后复核；发布状态见任务进度"])
    labels = {"completed": "写入并完成复核", "skipped": "最终无需再写", "hold": "仍挂账",
              "exception": "异常", "conflict": "冲突", "total": "核销记录数"}
    for name, label in labels.items():
        summary.append([label, payload["counts"].get(name, 0)])
    summary.append(["实际写入记录数", payload["write_count"]])

    initial = indexed_decisions(initial_result)
    checks = checked_records(checked_plan, {key for key, row in initial.items() if row["bucket"] == "auto"})
    if set(initial) != {row["case_id"] for row in payload["records"]}:
        workbook.close()
        raise ValueError("报告明细与首次判定记录不一致")
    codes = common.load_codes()
    details = workbook.create_sheet("核销明细", 1)
    headers = list(HEADERS)
    headers[1] = "首次处理"
    details.append([*headers, "实际执行", "最终状态", "最终原因", "首次校验及执行依据"])
    execution_labels = {"written_verified": "已写入并回读", "skipped": "首次校验跳过",
                        "conflict": "首次校验冲突", "not_executed": "未独立写入"}
    for record in payload["records"]:
        key = record["case_id"]
        item = checks["write"].get(key) or checks["skip"].get(key) or checks["conflict"].get(key) or initial[key]
        status = ("今天要填" if key in checks["write"] else "已填过·跳过" if key in checks["skip"]
                  else "冲突·需你定" if key in checks["conflict"]
                  else "挂账待办" if initial[key]["bucket"] == "hold" else "异常")
        details.append([*_row(item, status, codes, action_override=record["execution_reason"]),
                        execution_labels[record["execution_status"]], labels[record["final_status"]],
                        record["final_reason"], record["execution_reason"]])
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.data_type = "s"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24
    for column in (3, len(HEADERS) + 3, len(HEADERS) + 4):
        details.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 60
    workbook.active = 0
    workbook.save(target)
    workbook.close()
    workbook_finalize.finalize_static_report(target)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="生成首次、实际和最终状态一致的核销报告")
    parser.add_argument("--workspace", required=True)
    parser.add_argument("--checked", required=True)
    args = parser.parse_args(argv)
    workspace = Path(args.workspace).resolve(strict=True)
    checked = Path(args.checked).resolve(strict=True)
    if not checked.is_relative_to(workspace):
        raise ValueError("最终报告计划超出任务工作区")
    payload = build(workspace, checked)
    tag = payload["reconciliation_date"].replace("-", "")
    output = workspace / "04_产出"
    initial_result = json.loads((output / f"判定结果_{tag}.json").read_text(encoding="utf-8"))
    checked_plan = json.loads(checked.read_text(encoding="utf-8"))
    write_report(payload, output / f"核销日清_{tag}.xlsx", initial_result, checked_plan)
    (output / f"最终核销结果_{tag}.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8",
    )
    print(json.dumps({"counts": payload["counts"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
