"""Public classify -> validate -> OOXML write -> readback regression scenarios.

Synthetic workbooks only. Added for the approved change; not run under the
repository's diff-only verification restriction.
"""
import copy
import datetime as dt

import openpyxl
import pytest

import apply_to_copy as A
import baseline_receipts as B
import build_worklist as W
import classify_hexiao as C
import fallback_allocation_ledger as F
import validate_plan as V

SO, SOD = "SO26060786", "SOD_SYNTHETIC_60786"
HEADERS = ["部门", "销售人员", "客户名称", "单号", "新智云单号", "应收金额",
           "计提金额", "回款明细", "是否结账（是/否）", "收款时间", "收款方式", "实收金额", "差异"]


def workbook(path, *, historical=True, sibling=False, baseline=2934.88,
             so=SO, sod=SOD, historical_amount=5731.55):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "明细"
    ws.append(HEADERS)
    ws.append(["测试", "测试", "测试", "", so, baseline, None, None, "否", None, None, sod, None])
    if historical:
        ws.append(["测试", "测试", "测试", "", so, None, None, historical_amount, "是", "2026-09-01", "汇", sod, None])
    if sibling:
        ws.append(["测试", "测试", "测试", "", so, 100, None, None, "否", None, None, "SOD_OTHER", None])
    wb.create_sheet("保护页").append(["保留内容"])
    wb.save(path)
    wb.close()
    return path


def record(amount=1000, cumulative=6731.55, *, ar="AR_NEW", day=2, delivery=8590.11, sibling=False):
    deliveries = {SOD: delivery, **({"SOD_OTHER": 100} if sibling else {})}
    return {"ar": ar, "so": SO, "sod": SOD, "amount_orig": amount, "amount_local": amount,
            "deliver_local": delivery, "cumulative_received_local": cumulative,
            "itemized_cumulative_authoritative": True, "currency": "人民币CNY", "status": "手动核销",
            "hexiao_date": dt.date(2026, 9, day), "shoukuan_date": dt.date(2026, 9, day),
            "writeoff_sequence_key": [f"2026-09-{day:02d}", ar + "_DETAIL"],
            "all_sods": list(deliveries), "sod_delivery_local": deliveries,
            "so_delivery_local": sum(deliveries.values())}


def plan(path, records, state=None):
    ledger = C.LedgerIndex(path)
    ledger.baseline_receipt_state = state or {}
    classified = C.classify_records(records, ledger)
    checked = V.validate(classified, V.read_ledger_rows(path), path)
    checked["hexiao_date"] = records[0]["hexiao_date"].isoformat()
    return classified, checked


def write(source, target, checked):
    assert checked["conflict"] == []
    assert A.precheck_before_write(checked, checked["write"], source) == []
    A.write_plan(source, target, checked["write"])
    assert A.verify_written(target, checked["write"]) == []
    return V.read_ledger_rows(target)


def test_existing_blank_receipt_continues_without_changing_original_receivable(tmp_path):
    source = workbook(tmp_path / "source.xlsx")
    _, checked = plan(source, [record()])
    assert len(checked["write"]) == 1
    rows = write(source, tmp_path / "output.xlsx", checked)
    assert rows[2]["应收金额"] == 2934.88 and rows[2]["回款明细"] is None
    assert rows[2]["是否结账"] == "否"
    assert rows[3]["回款明细"] == 5731.55
    assert rows[4]["应收金额"] is None and rows[4]["回款明细"] == 1000
    assert all(row["计提"] is None for row in rows.values())
    saved = F.commit(tmp_path, checked)[0]
    assert saved.is_file()
    state = F.load(tmp_path)["baseline_receipts"]
    _, replay = plan(tmp_path / "output.xlsx", [record()], state)
    assert replay["write"] == [] and replay["conflict"] == [] and len(replay["skip"]) == 1


def test_final_receipt_settles_original_row_and_accrues_delivery_once(tmp_path):
    source = workbook(tmp_path / "source.xlsx")
    rec = record(2858.56, 8590.11)
    _, checked = plan(source, [rec])
    rows = write(source, tmp_path / "settled.xlsx", checked)
    assert all(row["是否结账"] == "是" for row in rows.values())
    assert sum(row["应收金额"] or 0 for row in rows.values()) == 2934.88
    assert round(sum(row["回款明细"] or 0 for row in rows.values()), 2) == 8590.11
    assert sum(row["计提"] or 0 for row in rows.values()) == 8590.11
    state = B.merge_journal({}, checked)
    _, replay = plan(tmp_path / "settled.xlsx", [rec], state)
    assert replay["write"] == [] and replay["conflict"] == [] and len(replay["skip"]) == 1


@pytest.mark.parametrize("amounts", [[5731.55, 2858.56], [2000, 3731.55, 2858.56]])
def test_same_day_chain_pins_mode_before_crossing_baseline_through_final_receipt(tmp_path, amounts):
    source = workbook(tmp_path / "source.xlsx", historical=False)
    cumulative, records = 0, []
    for index, amount in enumerate(amounts):
        cumulative = round(cumulative + amount, 2)
        records.append(record(amount, cumulative, ar=f"AR_{index}"))
    _, checked = plan(source, records)
    assert len(checked["write"]) == len(amounts)
    rows = write(source, tmp_path / "output.xlsx", checked)
    assert rows[2]["应收金额"] == 2934.88 and rows[2]["回款明细"] is None
    assert [rows[index + 3]["回款明细"] for index in range(len(amounts))] == amounts
    assert all(rows[index + 3]["应收金额"] is None for index in range(len(amounts)))
    assert sum(row["计提"] or 0 for row in rows.values()) == 8590.11


def test_same_day_tail_is_settled_by_last_event(tmp_path):
    source = workbook(tmp_path / "source.xlsx")
    _, checked = plan(source, [record(2858, 8589.55, ar="AR_A"), record(.56, 8590.11, ar="AR_B")])
    assert len(checked["write"]) == 2
    rows = write(source, tmp_path / "output.xlsx", checked)
    assert rows[4]["回款明细"] == 2858 and rows[4]["计提"] is None
    assert rows[5]["回款明细"] == .56 and rows[5]["计提"] == 8590.11


@pytest.mark.parametrize("amount", [2858.06, 2859.06])
def test_single_receipt_keeps_existing_one_yuan_settlement_tolerance(tmp_path, amount):
    source = workbook(tmp_path / "source.xlsx")
    _, checked = plan(source, [record(amount, round(5731.55 + amount, 2))])
    rows = write(source, tmp_path / "output.xlsx", checked)
    assert rows[4]["回款明细"] == amount and rows[4]["计提"] == 8590.11


def test_identical_signature_without_parent_identity_is_conflict(tmp_path):
    source = workbook(tmp_path / "source.xlsx")
    _, checked = plan(source, [record(5731.55, 5731.55, day=1)])
    assert checked["write"] == [] and checked["skip"] == [] and len(checked["conflict"]) == 1


def test_multiple_records_cannot_share_the_same_case_identity(tmp_path):
    source = workbook(tmp_path / "source.xlsx")
    first = record(1000, 6731.55)
    second = record(500, 7231.55)
    second["writeoff_sequence_key"] = ["2026-09-02", "ANOTHER_DETAIL"]
    _, checked = plan(source, [first, second])
    assert checked["write"] == [] and len(checked["conflict"]) == 2


def test_different_parents_with_same_amount_date_and_method_remain_distinct(tmp_path):
    source = workbook(tmp_path / "source.xlsx")
    _, first = plan(source, [record(1000, 6731.55, ar="AR_A")])
    target = tmp_path / "first.xlsx"
    write(source, target, first)
    state = B.merge_journal({}, first)
    _, second = plan(target, [record(1000, 7731.55, ar="AR_B")], state)
    assert len(second["write"]) == 1
    rows = write(target, tmp_path / "second.xlsx", second)
    assert rows[4]["回款明细"] == rows[5]["回款明细"] == 1000
    state = B.merge_journal(state, second)
    _, replay = plan(tmp_path / "second.xlsx", [record(1000, 7731.55, ar="AR_B")], state)
    assert len(replay["skip"]) == 1 and replay["write"] == []


def test_changed_or_partly_written_event_does_not_skip(tmp_path):
    source = workbook(tmp_path / "source.xlsx")
    _, checked = plan(source, [record()])
    target = tmp_path / "output.xlsx"
    write(source, target, checked)
    state = B.merge_journal({}, checked)
    _, changed = plan(target, [record(1100, 6831.55)], state)
    assert len(changed["conflict"]) == 1
    wb = openpyxl.load_workbook(target)
    wb["明细"].cell(2, 9, "是")
    wb.save(target)
    wb.close()
    _, partial = plan(target, [record()], state)
    assert len(partial["conflict"]) == 1 and partial["skip"] == []


def test_write_plan_cannot_inflate_amount_by_changing_only_candidate_fields(tmp_path):
    source = workbook(tmp_path / "source.xlsx")
    _, checked = plan(source, [record()])
    item = copy.deepcopy(checked["write"][0])
    step = item["row_operation"]["steps"][0]
    step.update(current_received=1100, cumulative_received=6831.55, remaining=1758.56)
    step["five_cols"]["回款明细"] = 1100
    item["five_cols"]["回款明细"] = 1100
    assert V.check_one(item, V.read_ledger_rows(source))["verdict"] == "conflict"


def test_final_plan_cannot_omit_accrual_without_so_gate(tmp_path):
    source = workbook(tmp_path / "source.xlsx")
    _, checked = plan(source, [record(2858.56, 8590.11)])
    item = copy.deepcopy(checked["write"][0])
    item["five_cols"]["计提"] = None
    item["derived_cols"] = {}
    item["row_operation"]["steps"][0]["five_cols"]["计提"] = None
    item["row_operation"]["steps"][0]["derived_cols"] = {}
    assert V.check_one(item, V.read_ledger_rows(source))["verdict"] == "conflict"


def test_paid_row_replay_does_not_release_accrual_for_unsettled_sod(tmp_path):
    source = workbook(tmp_path / "source.xlsx", sibling=True)
    rec = record(sibling=True)
    _, checked = plan(source, [rec])
    target = tmp_path / "partial.xlsx"
    write(source, target, checked)
    state = B.merge_journal({}, checked)
    sibling = {**record(100, 100, ar="AR_OTHER", sibling=True), "sod": "SOD_OTHER", "deliver_local": 100}
    classified, _ = plan(target, [rec, sibling], state)
    assert all(item["five_cols"].get("计提") is None for item in classified["auto"])
    assert all(not item.get("so_accrual_backfills") for item in classified["auto"])


def test_sod_settles_with_deferred_accrual_until_siblings_finish(tmp_path):
    source = workbook(tmp_path / "source.xlsx", sibling=True)
    rec = record(2858.56, 8590.11, sibling=True)
    _, checked = plan(source, [rec])
    rows = write(source, tmp_path / "output.xlsx", checked)
    assert all(row["计提"] is None for row in rows.values())
    assert all(row["是否结账"] == "是" for row in rows.values() if row["SOD"] == SOD)
    assert next(row for row in rows.values() if row["SOD"] == "SOD_OTHER")["是否结账"] == "否"


def test_delivery_equal_to_original_receivable_keeps_ordinary_mode(tmp_path):
    source = workbook(tmp_path / "source.xlsx", historical=False, baseline=8590.11)
    classified, checked = plan(source, [record(1000, 1000)])
    assert all(not item.get("baseline_receipt_audit") for item in classified["auto"])
    assert checked["write"][0]["row_operation"]["type"] == "split_below"


def test_later_sibling_settlement_backfills_accrual_and_publishes_replay_evidence(tmp_path):
    source = workbook(tmp_path / "source.xlsx", sibling=True)
    old_event = record(2858.56, 8590.11, sibling=True)
    _, initial = plan(source, [old_event])
    intermediate = tmp_path / "partial-so.xlsx"
    write(source, intermediate, initial)
    F.commit(tmp_path, initial)
    state = F.load(tmp_path)["baseline_receipts"]
    other = {**record(100, 100, ar="AR_OTHER", day=3, sibling=True), "sod": "SOD_OTHER", "deliver_local": 100}
    # Put the replay last: it must not become a writer of an already-paid row.
    _, final = plan(intermediate, [other, old_event], state)
    assert len(final["skip"]) == 1 and len(final["write"]) == 1
    target = tmp_path / "complete-so.xlsx"
    rows = write(intermediate, target, final)
    group = [row for row in rows.values() if row["SOD"] == SOD]
    assert sum(row["应收金额"] or 0 for row in group) == 2934.88
    assert round(sum(row["回款明细"] or 0 for row in group), 2) == 8590.11
    assert [row["计提"] for row in group].count(8590.11) == 1
    F.commit(tmp_path, final)
    state = F.load(tmp_path)["baseline_receipts"]
    assert state[B.group_key(SO, SOD)]["accrual"] == 8590.11
    _, replay = plan(target, [old_event], state)
    assert replay["write"] == [] and replay["conflict"] == [] and len(replay["skip"]) == 1


def test_deferred_plan_cannot_restore_accrual_while_sibling_is_open(tmp_path):
    source = workbook(tmp_path / "source.xlsx", sibling=True)
    _, checked = plan(source, [record(2858.56, 8590.11, sibling=True)])
    item = copy.deepcopy(checked["write"][0])
    item["five_cols"]["计提"] = 8590.11
    item["derived_cols"] = {"差异": -5655.23}
    item["row_operation"]["steps"][0]["five_cols"]["计提"] = 8590.11
    item["row_operation"]["steps"][0]["derived_cols"] = {"差异": -5655.23}
    assert V.check_one(item, V.read_ledger_rows(source))["verdict"] == "conflict"


def test_group_identity_and_source_rows_cannot_change_after_planning(tmp_path):
    source = workbook(tmp_path / "source.xlsx")
    _, checked = plan(source, [record()])
    item = copy.deepcopy(checked["write"][0])
    item["row_operation"]["group_key"] = B.group_key("SO_OTHER", SOD)
    assert V.check_one(item, V.read_ledger_rows(source))["verdict"] == "conflict"
    rows = V.read_ledger_rows(source)
    rows[2]["应收金额"] = 2935
    assert V.check_one(checked["write"][0], rows)["verdict"] == "conflict"


def test_report_exposes_mode_and_history_without_changing_sheet_name(tmp_path):
    source = workbook(tmp_path / "source.xlsx")
    classified, checked = plan(source, [record()])
    report = tmp_path / "report.xlsx"
    W.build_workbook(classified, checked, report)
    wb = openpyxl.load_workbook(report, data_only=True)
    assert "核销明细" in wb.sheetnames and "回款明细" not in wb.sheetnames
    sheet = wb["核销明细"]
    row = dict(zip([cell.value for cell in sheet[1]], [cell.value for cell in sheet[2]]))
    assert row["此前累计回款"] == 5731.55 and row["本次后剩余未收"] == 1858.56
    assert row["原始应收合计"] == 2934.88 and row["本次核销金额"] == 1000
    wb.close()


def group_payment(so, lines, amount, cumulative, *, ar="AR_GROUP", delivery=None):
    total = round(sum(lines.values()), 2) if delivery is None else delivery
    return {
        "ar": ar, "hexiao_date": dt.date(2026, 9, 2), "arrival_date": dt.date(2026, 9, 2),
        "amount_orig": amount, "amount_local": amount, "currency": "人民币CNY", "status": "核销成功",
        "orders": [{"so": so, "deliver": total, "currency": "人民币CNY"}],
        "writeoffs": {so: amount}, "writeoffs_local": {so: amount},
        "cumulative_writeoffs": {so: cumulative}, "cumulative_writeoffs_local": {so: cumulative},
        "sod_lines": {so: [{"sod": sod, "deliver": value} for sod, value in lines.items()]},
        "_writeoff_sequence_key_by_so": {so: ["2026-09-02", ar + "_DETAIL"]},
    }


@pytest.mark.parametrize("so,sod,baseline,history,lines", [
    ("SO26060786", "SOD26061011", 2934.88, 5731.55,
     {"SOD26061011": 2934.88, "SOD26080218": 5519.51, "SOD26080222": 135.72}),
    ("SO_GENERIC", "SOD_INITIAL", 100, 180,
     {"SOD_INITIAL": 100, "SOD_LATER_A": 130, "SOD_LATER_B": 70}),
])
def test_so_delivery_reaches_existing_group_before_sod_capacity_check(tmp_path, so, sod, baseline, history, lines):
    source = workbook(tmp_path / "source.xlsx", so=so, sod=sod, baseline=baseline, historical_amount=history)
    delivery = round(sum(lines.values()), 2)
    amount = round(delivery - history, 2)
    records = C.expand_payment(group_payment(so, lines, amount, delivery), {})
    classified, checked = plan(source, records)
    assert classified["hold"] == [] and len(checked["write"]) == 1
    item = checked["write"][0]
    assert item["split_payment_source"]["so_delivery_local"] == delivery
    assert item["baseline_receipt_audit"]["historical_received"] == history
    assert item["baseline_receipt_audit"]["remaining"] == 0
    assert item["so_accrual_audit"]["all_settled"] is True
    assert item.get("so_accrual_backfills", []) == []
    target = tmp_path / "final.xlsx"
    rows = write(source, target, checked)
    assert sum(row["应收金额"] or 0 for row in rows.values()) == baseline
    assert rows[2]["回款明细"] is None and rows[3]["回款明细"] == history
    assert rows[4]["回款明细"] == amount and rows[4]["应收金额"] is None
    assert all(row["是否结账"] == "是" for row in rows.values())
    assert sum(row["计提"] or 0 for row in rows.values()) == delivery
    F.commit(tmp_path, checked)
    _, replay = plan(target, records, F.load(tmp_path)["baseline_receipts"])
    assert len(replay["skip"]) == 1 and replay["write"] == [] and replay["conflict"] == []


def test_matching_source_sod_amount_still_uses_group_delivery_and_later_continues(tmp_path):
    source = workbook(tmp_path / "source.xlsx")
    lines = {SOD: 2934.88, "SOD_LATER_A": 5519.51, "SOD_LATER_B": 135.72}
    first_records = C.expand_payment(group_payment(SO, lines, 135.72, 5867.27, ar="AR_FIRST"), {})
    assert first_records[0]["sod"] == "SOD_LATER_B"  # Old subset-selection path.
    _, first = plan(source, first_records)
    middle = tmp_path / "middle.xlsx"
    rows = write(source, middle, first)
    assert rows[2]["是否结账"] == "否" and rows[4]["计提"] is None
    assert first["write"][0]["baseline_receipt_audit"]["remaining"] == 2722.84
    state = B.merge_journal({}, first)
    second_records = C.expand_payment(group_payment(SO, lines, 2722.84, 8590.11, ar="AR_SECOND"), {})
    _, second = plan(middle, second_records, state)
    rows = write(middle, tmp_path / "final.xlsx", second)
    assert rows[5]["计提"] == 8590.11 and rows[2]["应收金额"] == 2934.88


def test_all_source_sod_slices_of_one_receipt_are_written_once(tmp_path):
    source = workbook(tmp_path / "source.xlsx", historical=False)
    lines = {SOD: 2934.88, "SOD_LATER_A": 5519.51, "SOD_LATER_B": 135.72}
    records = C.expand_payment(group_payment(SO, lines, 8590.11, 8590.11), {})
    assert len(records) == 3
    _, checked = plan(source, records)
    assert len(checked["write"]) == 1
    rows = write(source, tmp_path / "output.xlsx", checked)
    assert len(rows) == 2 and rows[3]["回款明细"] == rows[3]["计提"] == 8590.11


def test_so_group_same_day_chain_uses_one_basis_before_crossing_original_receivable(tmp_path):
    source = workbook(tmp_path / "source.xlsx", historical=False)
    lines = {SOD: 2934.88, "SOD_LATER_A": 5519.51, "SOD_LATER_B": 135.72}
    records, cumulative = [], 0
    amounts = [2000, 3731.55, 2858.56]
    for index, amount in enumerate(amounts):
        cumulative = round(cumulative + amount, 2)
        records.extend(C.expand_payment(group_payment(SO, lines, amount, cumulative, ar=f"AR_{index}"), {}))
    _, checked = plan(source, records)
    assert len(checked["write"]) == 3
    rows = write(source, tmp_path / "output.xlsx", checked)
    assert rows[2]["应收金额"] == 2934.88 and rows[2]["回款明细"] is None
    assert [rows[index + 3]["回款明细"] for index in range(3)] == amounts
    assert [rows[index + 3]["计提"] for index in range(3)] == [None, None, 8590.11]


def test_so_group_never_combines_independent_ledger_sods(tmp_path):
    source = workbook(tmp_path / "source.xlsx", sibling=True)
    lines = {SOD: 2934.88, "SOD_OTHER": 100, "SOD_MISSING": 5555.23}
    classified, checked = plan(source, C.expand_payment(group_payment(SO, lines, 2858.56, 8590.11), {}))
    assert checked["write"] == [] and classified["hold"]
    assert all(not item.get("baseline_receipt_audit") for item in classified["hold"])


def test_so_latest_delivery_is_not_replaced_by_source_detail_amount_sum(tmp_path):
    source = workbook(tmp_path / "source.xlsx")
    lines = {SOD: 2934.88, "SOD_LATER": 5655.23}
    records = C.expand_payment(group_payment(SO, lines, 3268.45, 9000, delivery=9000), {})
    _, checked = plan(source, records)
    rows = write(source, tmp_path / "output.xlsx", checked)
    assert rows[4]["计提"] == 9000 and rows[2]["应收金额"] == 2934.88


def test_single_source_sod_also_uses_latest_so_delivery(tmp_path):
    source = workbook(tmp_path / "source.xlsx")
    records = C.expand_payment(group_payment(SO, {SOD: 2934.88}, 3268.45, 9000, delivery=9000), {})
    _, checked = plan(source, records)
    rows = write(source, tmp_path / "output.xlsx", checked)
    assert rows[4]["计提"] == 9000 and rows[2]["应收金额"] == 2934.88


def test_so_group_rejects_missing_delivery_instead_of_using_receivable(tmp_path):
    source = workbook(tmp_path / "source.xlsx")
    lines = {SOD: 2934.88, "SOD_LATER": 5655.23}
    payment = group_payment(SO, lines, 2858.56, 8590.11)
    payment["orders"][0]["deliver"] = None
    records = C.expand_payment(payment, {})
    classified, checked = plan(source, records)
    assert checked["write"] == [] and classified["hold"]
    assert "SO 最新交付额" in classified["hold"][0]["reason"]


def test_so_group_write_precheck_binds_source_scope_and_all_so_rows(tmp_path):
    source = workbook(tmp_path / "source.xlsx")
    lines = {SOD: 2934.88, "SOD_LATER": 5655.23}
    _, checked = plan(source, C.expand_payment(group_payment(SO, lines, 2858.56, 8590.11), {}))
    item = copy.deepcopy(checked["write"][0])
    item["split_payment_source"]["so_delivery_local"] = 2934.88
    assert V.check_one(item, V.read_ledger_rows(source))["verdict"] == "conflict"
    rows = V.read_ledger_rows(source)
    rows[5] = {**rows[2], "SOD": "SOD_LATER", "应收金额": 10}
    assert V.check_one(checked["write"][0], rows)["verdict"] == "conflict"


def test_equal_so_delivery_uses_whole_group_for_ordinary_partial_and_final(tmp_path):
    source = workbook(tmp_path / "source.xlsx", historical=False, baseline=300)
    lines = {SOD: 100, "SOD_LATER": 200}
    _, first = plan(source, C.expand_payment(group_payment(SO, lines, 100, 100, ar="AR_FIRST"), {}))
    assert first["write"][0]["row_operation"]["type"] == "split_below"
    assert not first["write"][0].get("baseline_receipt_audit")
    middle = tmp_path / "middle.xlsx"
    rows = write(source, middle, first)
    assert any(row["是否结账"] != "是" for row in rows.values())
    assert sum(row["应收金额"] or 0 for row in rows.values()) == 300
    assert all(row["计提"] is None for row in rows.values())
    state = B.merge_journal({}, first)
    assert state[B.group_key(SO, SOD)]["scope_only"] is True
    _, final = plan(middle, C.expand_payment(group_payment(SO, lines, 200, 300, ar="AR_FINAL"), {}), state)
    rows = write(middle, tmp_path / "final.xlsx", final)
    assert all(row["是否结账"] == "是" for row in rows.values())
    assert sum(row["计提"] or 0 for row in rows.values()) == 300
    assert sum(row["回款明细"] or 0 for row in rows.values()) == 300


@pytest.mark.parametrize("change", ["source_shrinks", "metadata_missing", "independent_group"])
@pytest.mark.parametrize("baseline", [3000, 9000])
def test_published_so_scope_never_falls_back_to_source_detail_amount(tmp_path, change, baseline):
    source = workbook(tmp_path / "source.xlsx", historical=False, baseline=baseline)
    lines = {SOD: 8000, "SOD_LATER": 1000}
    _, first = plan(source, C.expand_payment(group_payment(SO, lines, 6000, 6000, ar="AR_FIRST"), {}))
    middle = tmp_path / "middle.xlsx"
    write(source, middle, first)
    state = B.merge_journal({}, first)
    if change == "source_shrinks":
        lines = {SOD: 8000}
    records = C.expand_payment(group_payment(SO, lines, 2000, 8000, delivery=9000, ar="AR_NEXT"), {})
    if change == "metadata_missing":
        for rec in records:
            rec.pop("so_receipt_source", None)
    if change == "independent_group":
        wb = openpyxl.load_workbook(middle)
        wb["明细"].append(["测试", "测试", "测试", "", SO, 1000, None, None, "否", None, None, "SOD_LATER", None])
        wb.save(middle)
        wb.close()
    classified, checked = plan(middle, records, state)
    assert checked["write"] == []
    assert classified["hold"] or checked["conflict"]


def test_ordinary_scope_is_checked_before_settled_skip(tmp_path):
    source = workbook(tmp_path / "source.xlsx", historical=False, baseline=300)
    classified, _ = plan(source, C.expand_payment(group_payment(SO, {SOD: 100, "SOD_LATER": 200}, 100, 100), {}))
    rows = V.read_ledger_rows(source)
    rows[2]["是否结账"] = "是"
    rows[3] = {**rows[2], "SOD": "SOD_LATER", "应收金额": 200, "是否结账": "否"}
    checked = V.validate(classified, rows)
    assert checked["write"] == [] and checked["skip"] == [] and checked["conflict"]


def test_missing_source_detail_amount_does_not_replace_known_so_delivery(tmp_path):
    source = workbook(tmp_path / "source.xlsx")
    payment = group_payment(SO, {SOD: 2934.88, "SOD_LATER": 5655.23}, 2858.56, 8590.11)
    payment["sod_lines"][SO][1]["deliver"] = None
    _, checked = plan(source, C.expand_payment(payment, {}))
    assert len(checked["write"]) == 1
    assert checked["write"][0]["baseline_receipt_audit"]["remaining"] == 0
