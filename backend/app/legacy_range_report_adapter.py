from __future__ import annotations

import argparse
import datetime as dt
import importlib.util
import json
import re
import sys
from copy import copy
from pathlib import Path
from types import ModuleType

import openpyxl
from openpyxl.styles import Font, PatternFill

CLEANUP_REPORT_PREFIXES = ("核销日清", "变更清单", "订单写入差异")
AGGREGATED_SHEETS = {
    "今日清单": "核销明细",
    "跨月计提补填": "跨月计提补填",
    "流转表怎么填": "流转表怎么填",
}

def _load_legacy_builder(script_path: Path) -> ModuleType:
    script_path = script_path.resolve()
    spec = importlib.util.spec_from_file_location("legacy_build_task_reports", script_path)
    if spec is None or spec.loader is None:
        raise RuntimeError("旧版范围报告脚本无法加载。")
    module = importlib.util.module_from_spec(spec)
    sys.path.insert(0, str(script_path.parent))
    try:
        spec.loader.exec_module(module)
    finally:
        sys.path.pop(0)
    return module


def _report_name(selected_days: list[dt.date]) -> str:
    start = selected_days[0].strftime("%Y%m%d")
    end = selected_days[-1].strftime("%Y%m%d")
    consecutive = len(selected_days) == (selected_days[-1] - selected_days[0]).days + 1
    if consecutive:
        return f"核销日清_{start}_{end}.xlsx"
    return f"核销日清_已选{len(selected_days)}日_{start}_{end}.xlsx"


def _daily_counts(output_dir: Path, day: dt.date) -> dict[str, object]:
    result_path = output_dir / f"判定结果_{day.strftime('%Y%m%d')}.json"
    if not result_path.is_file():
        return {}
    payload = json.loads(result_path.read_text(encoding="utf-8"))
    counts = payload.get("counts") or {}
    return {
        "到账笔数": payload.get("payment_count", ""),
        "订单行数": counts.get("total", ""),
        "自动": counts.get("auto", ""),
        "挂账": counts.get("hold", ""),
        "异常": counts.get("exception", ""),
    }


def _daily_report_files(output_dir: Path, prefix: str, day: dt.date) -> list[Path]:
    token = day.strftime("%Y%m%d")
    pattern = re.compile(rf"^{re.escape(prefix)}_{token}(?:_\d{{4}})?\.xlsx$")
    return sorted(
        path
        for path in output_dir.glob(f"{prefix}_{token}*.xlsx")
        if pattern.fullmatch(path.name)
    )


def _append_sheet(target, source, day: dt.date, header_written: bool) -> bool:
    rows = list(source.iter_rows())
    if not rows:
        return header_written
    if not header_written:
        target.append(["核销日期", *(cell.value for cell in rows[0])])
        for column, cell in enumerate(rows[0], start=2):
            output = target.cell(row=1, column=column)
            if cell.has_style:
                output._style = copy(cell._style)
            output.number_format = cell.number_format
            output.alignment = copy(cell.alignment)
            output.protection = copy(cell.protection)
        target.cell(row=1, column=1).font = Font(bold=True)
        target.cell(row=1, column=1).fill = PatternFill("solid", fgColor="D9EAF7")
        header_written = True
    for source_row in rows[1:]:
        values = [cell.value for cell in source_row]
        if not any(value not in (None, "") for value in values):
            continue
        target.append([day.isoformat(), *values])
        for column, cell in enumerate(source_row, start=2):
            output = target.cell(row=target.max_row, column=column)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                output.data_type = "s"
            if cell.has_style:
                output._style = copy(cell._style)
            output.number_format = cell.number_format
            output.alignment = copy(cell.alignment)
            output.protection = copy(cell.protection)
    return header_written


def build_selected_reports(
    script_path: Path,
    workspace: Path,
    selected_dates: list[str],
    *,
    empty_dates: list[str] | None = None,
) -> Path:
    selected_days = sorted({dt.date.fromisoformat(value) for value in selected_dates})
    if not selected_days:
        raise ValueError("至少需要一个核销日期。")
    empty_days = {dt.date.fromisoformat(value) for value in (empty_dates or [])}
    if empty_days - set(selected_days):
        raise ValueError("空核销日期必须属于本批已选日期。")
    module = _load_legacy_builder(script_path)
    finalize = getattr(getattr(module, "workbook_finalize", None), "finalize_static_report", None)
    if not callable(finalize):
        raise RuntimeError("旧版范围报告脚本不支持平台兼容调用。")
    output_dir = workspace.resolve() / "04_产出"
    daily_sources: set[Path] = set()
    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "任务范围"
    summary.append(["核销日期", "日报状态", "到账笔数", "订单行数", "自动", "挂账", "异常"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    aggregated: dict[str, object] = {}
    header_written: set[str] = set()
    for day in selected_days:
        daily = output_dir / f"核销日清_{day.strftime('%Y%m%d')}.xlsx"
        for prefix in CLEANUP_REPORT_PREFIXES:
            daily_sources.update(_daily_report_files(output_dir, prefix, day))
        counts = _daily_counts(output_dir, day)
        if daily.is_file():
            status = "已纳入"
        elif day in empty_days:
            status = "无核销记录，已跳过"
        else:
            raise FileNotFoundError("缺少日期级核销日清。")
        summary.append([
            day.isoformat(), status,
            counts.get("到账笔数", ""), counts.get("订单行数", ""),
            counts.get("自动", ""), counts.get("挂账", ""), counts.get("异常", ""),
        ])
        if not daily.is_file():
            continue
        source_workbook = openpyxl.load_workbook(daily, data_only=False, read_only=False)
        try:
            for source_title, target_title in AGGREGATED_SHEETS.items():
                if source_title not in source_workbook.sheetnames:
                    continue
                target = aggregated.get(target_title)
                if target is None:
                    target = workbook.create_sheet(target_title)
                    aggregated[target_title] = target
                if _append_sheet(
                    target,
                    source_workbook[source_title],
                    day,
                    target_title in header_written,
                ):
                    header_written.add(target_title)
        finally:
            source_workbook.close()
    summary.freeze_panes = "A2"
    summary.auto_filter.ref = summary.dimensions
    for column, width in zip("ABCDEFG", (14, 16, 12, 12, 10, 10, 10), strict=True):
        summary.column_dimensions[column].width = width
    for sheet in aggregated.values():
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.column_dimensions["A"].width = 14
    report = output_dir / _report_name(selected_days)
    workbook.save(report)
    finalize(report)
    for source in sorted(daily_sources):
        if source.resolve() != report.resolve() and source.is_file():
            source.unlink()
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="兼容旧版 Skill 快照生成已选日期范围报告")
    parser.add_argument("--script", required=True)
    parser.add_argument("--workspace", required=True)
    parser.add_argument("--date", action="append", required=True)
    parser.add_argument("--empty-date", action="append", default=[])
    args = parser.parse_args(argv)
    try:
        report = build_selected_reports(
            Path(args.script),
            Path(args.workspace),
            args.date,
            empty_dates=args.empty_date,
        )
    except Exception:
        print("ERROR: 旧版范围报告生成失败。", file=sys.stderr)
        return 2
    print(report)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
