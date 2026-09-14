"""Recheck settled skips with the full, previously checked allocation component."""
import copy
import sys
import tempfile
import unittest
from pathlib import Path
import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'vendor/scripts'))
import validate_plan as V
import apply_all as A
import common

class SkipSequenceRecheckTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.path=Path(self.temp.name)/'ledger.xlsx'
        book=openpyxl.Workbook();sheet=book.active;sheet.title='明细'
        sheet.append(['SO','SOD','应收金额','计提','回款明细','是否结账','收款时间','收款方式'])
        sheet.append(['SO_DONE','SOD_DONE',100,0,100,'是','2026-01-01','汇款'])
        book.save(self.path);book.close()
        self.skip=dict(case_id='AR_A|SO_DONE',ar='AR_A',so='SO_DONE',sod='SOD_DONE',
            code='OK_SO_ALREADY_SETTLED',ledger_year=2026,ledger_row_ref=2,
            fallback_batch_cases=['AR_A|SO_DONE','AR_A|SO_OPEN'],_check={'verdict':'skip'})
        self.other=dict(case_id='AR_A|SO_OPEN',ar='AR_A',so='SO_OPEN',ledger_year=2026,
            fallback_batch_cases=['AR_A|SO_DONE','AR_A|SO_OPEN'],_check={'verdict':'write'})
        self.plan=dict(write=[self.other],skip=[self.skip],conflict=[],
            ledger_path=str(self.path),ledger_sha256=common.sha256_file(self.path),
            ledger_checks={'2026':{'path':str(self.path),'sha256':common.sha256_file(self.path)}})

    def test_checked_write_dependency_is_visible_during_skip_recheck(self):
        self.assertEqual(V.recheck_so_skips(self.plan,self.path),[])

    def test_year_partition_keeps_checked_dependency_from_another_year(self):
        self.other['ledger_year']=2025
        subplan=A._year_subplan(self.plan,2026,self.path)
        self.assertEqual(V.recheck_so_skips(subplan,self.path),[])

    def test_missing_dependency_still_blocks(self):
        self.plan['write']=[]
        self.assertTrue(V.recheck_so_skips(self.plan,self.path))

    def test_conflicting_dependency_still_blocks(self):
        self.plan['write']=[];self.plan['conflict']=[self.other]
        self.other['_check']={'verdict':'conflict','reason':'unresolved'}
        self.assertTrue(V.recheck_so_skips(self.plan,self.path))

    def test_unchecked_dependency_still_blocks(self):
        self.other.pop('_check')
        self.assertTrue(V.recheck_so_skips(self.plan,self.path))

    def test_changed_workbook_still_blocks(self):
        book=openpyxl.load_workbook(self.path);book.active.cell(2,6,'否');book.save(self.path);book.close()
        self.assertTrue(V.recheck_so_skips(self.plan,self.path))

    def test_no_longer_settled_row_still_blocks_with_matching_hash(self):
        book=openpyxl.load_workbook(self.path);book.active.cell(2,6,'否');book.save(self.path);book.close()
        self.plan['ledger_sha256']=common.sha256_file(self.path)
        self.assertTrue(V.recheck_so_skips(self.plan,self.path))

    def test_recheck_does_not_mutate_checked_plan(self):
        before=copy.deepcopy(self.plan)
        V.recheck_so_skips(self.plan,self.path)
        self.assertEqual(self.plan,before)

    def test_conflicting_other_year_is_not_lost(self):
        self.other['ledger_year']=2025;self.other['_check']={'verdict':'conflict'}
        self.plan['write']=[];self.plan['conflict']=[self.other]
        self.assertTrue(V.recheck_so_skips(A._year_subplan(self.plan,2026,self.path),self.path))
