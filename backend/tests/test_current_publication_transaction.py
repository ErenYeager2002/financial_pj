"""Real publication/savepoint/copy cleanup from an already verified stage.

Upstream classification and stage construction are fixture inputs; publication,
material replacement, readback, owner checks and transaction fencing are real.
"""
import hashlib,json
from datetime import datetime,UTC,timedelta
from pathlib import Path
import openpyxl
import pytest
from sqlalchemy import select
from sqlalchemy.orm import Session
from app import ar_execution_runner as runner,workflow_service as service
from app.models import WorkflowSession,WorkflowAction,WorkflowMaterialSet,FileRecord
from test_refactor_event_transactions import database
from test_publication_current_evidence import seed


@pytest.mark.parametrize('fault',[None,'after_material_creation','partial_artifact_copy','workbook_changed','final_report_changed','complete_chain'])
def test_verified_stage_publication_is_atomic_with_file_cleanup(database,tmp_path,monkeypatch,fault):
    for mod in (runner,service):monkeypatch.setattr(mod,'workflow_root',lambda *_:tmp_path)
    monkeypatch.setattr(service,'_workflow_storage_root',lambda *_:tmp_path)
    with Session(database) as db:
        seed(db,tmp_path)
        flow=db.get(WorkflowSession,'cancel-workflow');flow.state='running'
        action=WorkflowAction(id='b'*32,workflow_id=flow.id,name='ar_publish_reconciliation',
            state='running',worker_id='test-worker',attempt_count=1,
            lease_expires_at=datetime.now(UTC)+timedelta(minutes=10))
        db.add(action);db.commit()
        workspace=tmp_path/'work';stage=workspace/'stage';books=stage/'02_我的表副本';books.mkdir(parents=True)
        outputs=stage/'04_产出';outputs.mkdir()
        for name in ('2026年盈亏.xlsx','到账流转.xlsx'):
            wb=openpyxl.Workbook();wb.active['A1']='verified synthetic';wb.save(books/name);wb.close()
        report=outputs/'核销日清_20260920.xlsx'
        wb=openpyxl.Workbook();wb.active['A1']='report';wb.save(report);wb.close()
        final=outputs/'final.json';final.write_text('{}')
        original={r.id:Path(r.stored_path).read_bytes() for r in db.scalars(select(FileRecord))}
        obj=object.__new__(runner.ArExecution)
        obj.db,obj.action,obj.workflow,obj.service=db,action,flow,service
        obj.workspace=workspace;obj.date=flow.reconciliation_date;obj.tag='20260920'
        obj.ledgers={2026:workspace/'02_我的表副本/2026年盈亏.xlsx'}
        obj.context={'flow_file':str(workspace/'02_我的表副本/到账流转.xlsx')}
        obj.execution={'publication':'not_published','material_set_id':flow.material_set_id,'material_version':1,
                       'reconciliation_date':flow.reconciliation_date,'skill_hash':flow.skill_hash,
                       'reconciliation_policy':'current-workbook-v1','steps':{}}
        fingerprints=obj._workbook_fingerprints(stage)
        obj.execution['steps']={'review_final_report':{'files':fingerprints},
            'build_final_report':{'files':fingerprints,'final_result':{'path':str(final),'fingerprint':hashlib.sha256(final.read_bytes()).hexdigest()}},
            'write_receipt_flow':{'flow_written':True}}
        obj.staging=lambda:(stage,stage/'checked.json')
        publish=service._publish_verified_material_set
        def fail_after_creation(*args,**kwargs):
            material,bindings=publish(*args,**kwargs)
            assert material.id!='published' and material.state=='current'
            raise RuntimeError('injected after material creation')
        error_type,error_match=RuntimeError,'injected after material creation'
        if fault=='after_material_creation':
            monkeypatch.setattr(service,'_publish_verified_material_set',fail_after_creation)
        elif fault=='partial_artifact_copy':
            register=service._register_artifact
            copied=[]
            def fail_second_copy(*args,**kwargs):
                if copied:raise RuntimeError('injected after first artifact copy')
                artifact=register(*args,**kwargs);copied.append(artifact)
                return artifact
            monkeypatch.setattr(service,'_register_artifact',fail_second_copy)
            error_match='injected after first artifact copy'
        elif fault=='workbook_changed':
            path=books/'2026年盈亏.xlsx'
            wb=openpyxl.load_workbook(path);wb.active['B9']='unexpected unrelated edit';wb.save(path);wb.close()
            error_type,error_match=ValueError,'暂存'
        elif fault=='final_report_changed':
            final.write_text('{"unexpected":"changed after review"}')
            error_type,error_match=ValueError,'最终核销结果与已复核版本不一致'
        if fault not in (None,'complete_chain'):
            with pytest.raises(error_type,match=error_match):
                obj.publish_reconciliation()
            db.rollback();db.expire_all()
            assert {r.id for r in db.scalars(select(FileRecord))}=={'annual','receipt'}
            assert [r.id for r in db.scalars(select(WorkflowMaterialSet))]==['published']
            assert db.get(WorkflowMaterialSet,'published').state=='current'
            assert not (tmp_path/'outputs'/action.id).exists()
        else:
            if fault=='complete_chain':
                from app.ar_execution_contract import PHASES,CONTRACT_VERSION
                from app.ar_execution_safety import register_effect_intent
                obj.execution.update(schema_version=CONTRACT_VERSION,
                    completed=[p.name for p in PHASES if p.name not in ('publish_reconciliation','complete_reconciliation')])
                obj.execution['steps']['write_ledger']={'ledger_written':True}
                checked=stage/'checked.json'
                checked.write_text(json.dumps({'hexiao_date':obj.date,'write':[],'skip':[]}))
                obj.context.update(workspace=str(workspace),plan_fingerprint=hashlib.sha256(checked.read_bytes()).hexdigest(),ar_execution=obj.execution)
                register_effect_intent(obj.context,flow,action,'publish_reconciliation')
                flow.context_json=json.dumps(obj.context);db.commit()
            result=obj.publish_reconciliation()
            if fault=='complete_chain':
                obj.execution['steps']['publish_reconciliation']=dict(result)
                obj.execution['completed'].append('publish_reconciliation')
                obj.execution['publication']='verified'
                result['ar_execution']=obj.execution
                runner.transition_phase(db,action,flow,result)
            db.commit();db.expire_all()
            material=db.get(WorkflowMaterialSet,result['material_set_id'])
            assert material.version==2 and material.parent_set_id=='published'
            assert material.state=='current' and db.get(WorkflowMaterialSet,'published').state=='superseded'
            assert len(result['artifacts'])==4
            assert len(list(db.scalars(select(FileRecord))))==6
            if fault=='complete_chain':
                complete_published_stage(db,obj,tmp_path,monkeypatch)
                assert len(list(db.scalars(select(FileRecord))))==7
        assert all(Path(db.get(FileRecord,key).stored_path).read_bytes()==raw for key,raw in original.items())


@pytest.mark.parametrize('mode',['unchanged','changed','symlink','wrong_action','referenced','unrelated'])
def test_rollback_cleanup_requires_exact_copy_receipt(database,tmp_path,monkeypatch,mode):
    monkeypatch.setattr(service,'workflow_root',lambda *_:tmp_path)
    with Session(database) as db:
        seed(db,tmp_path)
        flow=db.get(WorkflowSession,'cancel-workflow')
        action_id='c'*32
        delivery=tmp_path/'outputs'/action_id;delivery.mkdir(parents=True)
        target=delivery/'result.json';raw=b'owned result';target.write_bytes(raw)
        receipt=dict(file_id='rolled-back-file',name=target.name,action_id=action_id,
                     size_bytes=len(raw),sha256=hashlib.sha256(raw).hexdigest())
        other=delivery/'keep.txt'
        if mode=='changed':target.write_bytes(b'newer result')
        if mode=='symlink':
            other.write_bytes(raw);target.unlink();target.symlink_to(other)
        if mode=='wrong_action':receipt['action_id']='another-action'
        if mode=='unrelated':other.write_bytes(b'unrelated')
        if mode=='referenced':
            db.add(FileRecord(id='other-reference',owner_id=flow.owner_id,department_id=flow.department_id,
                skill_id=flow.skill_id,workflow_id=flow.id,kind='output',original_name=target.name,
                stored_path=str(target.resolve()),size_bytes=len(raw),sha256=receipt['sha256']))
            db.flush()
        if mode in ('changed','symlink','wrong_action'):
            with pytest.raises(ValueError):service._discard_registered_artifacts(db,flow,[receipt],action_id)
            assert target.exists()
        else:
            service._discard_registered_artifacts(db,flow,[receipt],action_id)
            assert target.exists()==(mode=='referenced')
            if mode=='unrelated':assert other.read_bytes()==b'unrelated'
            if mode=='unchanged':
                assert not delivery.exists()
                service._discard_registered_artifacts(db,flow,[receipt],action_id)
        assert (tmp_path/'annual.xlsx').read_bytes()==b'synthetic annual'
        assert (tmp_path/'receipt.xlsx').read_bytes()==b'synthetic receipt'


def complete_published_stage(db,obj,root,monkeypatch,*,prepare_synthetic_journals=True,expected_parents=None):
    """Continue the actual queued completion with the real lab script and journal."""
    import shutil,subprocess,sys
    from app import ar_publication as publication,ar_formal_ledger_service as formal,ar_process_evidence as evidence
    from app.ar_execution_safety import register_effect_intent
    for module in (publication,formal,evidence):
        monkeypatch.setattr(module,'workflow_root',lambda *_:root)
    flow=obj.workflow
    action=db.scalar(select(WorkflowAction).where(WorkflowAction.workflow_id==flow.id,
                                                WorkflowAction.name=='ar_complete_reconciliation'))
    assert action is not None and action.state=='queued'
    action.state='running';action.worker_id='test-worker';action.attempt_count=1
    action.lease_expires_at=datetime.now(UTC)+timedelta(minutes=10)
    obj.action=action
    obj.context=json.loads(flow.context_json);obj.execution=obj.context['ar_execution']
    register_effect_intent(obj.context,flow,action,'complete_reconciliation')
    flow.context_json=json.dumps(obj.context);db.commit()
    db.info.pop('ar_execution_lock',None)
    vendor=Path(__file__).resolve().parents[2]/'skills/ar-hexiao-daily-lab/vendor'
    if not (root/'skill/vendor').exists():
        shutil.copytree(vendor,root/'skill/vendor',ignore=shutil.ignore_patterns('__pycache__'))
    obj.scripts=root/'skill/vendor/scripts'
    stage,_=obj.staging()
    if prepare_synthetic_journals:
        subprocess.run([sys.executable,'-B','-c',
            'from pathlib import Path; import sys; import rescan_holds as H; w=Path(sys.argv[1]); H.save_ledger(H.ledger_path(w),[])',
            str(stage)],cwd=obj.scripts,check=True,capture_output=True)
        journal=stage/'03_台账'
        (journal/'父回款顺序分配台账.json').write_bytes(b'invalid old audit only')
    result=obj.complete_reconciliation()
    obj.execution['steps']['complete_reconciliation']=dict(result)
    obj.execution['completed'].append('complete_reconciliation')
    result.update(ar_execution=obj.execution,process_evidence_version=evidence.SCHEMA_VERSION,
                  process_records=list(action._ar_process_records))
    assert action._ar_process_exit_confirmed
    assert len(result['process_records'])==1
    # Only constructor/snapshot setup is supplied; all method checks run unchanged.
    monkeypatch.setattr(runner,'ArExecution',lambda *_:obj)
    runner.transition_phase(db,action,flow,result);db.commit();db.expire_all()
    assert flow.state=='succeeded' and action.state=='succeeded'
    assert all(a.state=='succeeded' for a in db.scalars(select(WorkflowAction)))
    payload,record,contents=formal.read_formal_ledger_bundle(db,flow)
    assert payload['publication']['material_set_id']==flow.material_set_id
    assert payload['publication']['material_version']==2
    def normalize(parents):
        return {ar:{k:v for k,v in entry.items() if k not in ('applied_at','last_verified_at','reused_successful_allocation')} for ar,entry in parents.items()}
    assert normalize(json.loads(contents['父回款顺序分配台账.json'])['parents'])==normalize(expected_parents or {})
    assert record.sha256==hashlib.sha256(Path(result['formal_ledger_candidate']['path']).read_bytes()).hexdigest()
