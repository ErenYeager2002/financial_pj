from __future__ import annotations

import json
import time
import uuid
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.auth import UserContext
from app.auth_models import User
from app.database import SessionLocal, engine
from app.models import FileRecord, RunRecord
from app.registry import registry
from app.resource_policy import upload_root
from app.run_service import confirm_run, create_run
from app.schemas import RunCreate
from app.settings import settings
from app.storage import sha256_file


IDEMPOTENCY_KEY = "stage10-postgres-pilot-v1"


def _write_workbook(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def _input_file(
    db,
    user: UserContext,
    name: str,
    headers: list[str],
    rows: list[list[object]],
) -> FileRecord:
    file_id = str(uuid.uuid4())
    folder = upload_root(user.user_id, file_id)
    folder.mkdir(parents=True, exist_ok=False)
    target = folder / name
    _write_workbook(target, headers, rows)
    record = FileRecord(
        id=file_id,
        owner_id=user.user_id,
        department_id=user.department_id,
        kind="input",
        original_name=name,
        stored_path=str(target.resolve()),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        size_bytes=target.stat().st_size,
        sha256=sha256_file(target),
    )
    db.add(record)
    db.flush()
    return record


def _wait_for_run(run_id: str, timeout_seconds: int = 90) -> RunRecord:
    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        with SessionLocal() as db:
            run = db.get(RunRecord, run_id)
            if run and run.state in {"succeeded", "failed", "timed_out", "cancelled"}:
                db.expunge(run)
                return run
        time.sleep(1)
    raise TimeoutError("合成任务未在限定时间内结束。")


def main() -> int:
    if engine.dialect.name != "postgresql":
        raise RuntimeError("Stage 10 试运行必须使用 PostgreSQL。")
    registry.refresh()
    if registry.errors:
        raise RuntimeError("Skill 注册表存在错误。")

    with SessionLocal() as db:
        admin = db.scalar(
            select(User)
            .where(User.role == "skill_admin", User.status == "active")
            .order_by(User.created_at)
        )
        if not admin:
            raise RuntimeError("没有可用于内部合成试运行的平台管理员。")
        user = UserContext(
            user_id=admin.id,
            display_name="Stage 10 合成试运行",
            role=admin.role,
            department_id=admin.department_id,
            username="stage10-pilot",
            auth_provider="internal-pilot",
        )
        existing = db.scalar(
            select(RunRecord).where(
                RunRecord.owner_id == user.user_id,
                RunRecord.idempotency_key == IDEMPOTENCY_KEY,
            )
        )
        if existing:
            run_id = existing.id
        else:
            bank = _input_file(
                db,
                user,
                "stage10_合成银行流水.xlsx",
                ["交易日期", "交易金额", "流水号"],
                [
                    ["2026-08-01", 100.0, "PILOT-B-001"],
                    ["2026-08-02", 200.0, "PILOT-B-002"],
                    ["2026-08-10", 999.0, "PILOT-B-003"],
                ],
            )
            ledger = _input_file(
                db,
                user,
                "stage10_合成财务总账.xlsx",
                ["凭证日期", "金额", "凭证号"],
                [
                    ["2026-08-01", 100.0, "PILOT-L-001"],
                    ["2026-08-03", 200.5, "PILOT-L-002"],
                    ["2026-08-20", 700.0, "PILOT-L-003"],
                ],
            )
            run = create_run(
                db,
                RunCreate(
                    skill_id="reconcile-bank",
                    message="Stage 10 PostgreSQL 合成试运行",
                    parameters={"amount_tolerance": 1, "date_tolerance_days": 2},
                    files={"bank_file": bank.id, "ledger_file": ledger.id},
                    idempotency_key=IDEMPOTENCY_KEY,
                ),
                user,
            )
            confirm_run(db, run, user)
            db.commit()
            run_id = run.id

    completed = _wait_for_run(run_id)
    if completed.state != "succeeded":
        raise RuntimeError(f"合成任务执行失败：{completed.state} {completed.error_message}")
    result = json.loads(completed.result_json or "{}")
    expected_summary = {
        "bank_records": 3,
        "ledger_records": 3,
        "matched": 2,
        "unmatched_bank": 1,
        "unmatched_ledger": 1,
    }
    if result.get("summary") != expected_summary:
        raise RuntimeError("合成任务结果摘要不符合预期。")

    output_checks: list[dict[str, object]] = []
    with SessionLocal() as db:
        outputs = list(
            db.scalars(
                select(FileRecord).where(
                    FileRecord.run_id == run_id,
                    FileRecord.kind == "output",
                )
            ).all()
        )
        if not outputs:
            raise RuntimeError("合成任务没有生成输出文件。")
        for item in outputs:
            path = Path(item.stored_path)
            if not path.is_file() or sha256_file(path) != item.sha256:
                raise RuntimeError("合成任务输出文件哈希校验失败。")
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet_names = workbook.sheetnames
            workbook.close()
            output_checks.append(
                {
                    "sha256": item.sha256,
                    "size_bytes": item.size_bytes,
                    "sheets": sheet_names,
                }
            )

    report = {
        "verified": True,
        "created_at": datetime.now(UTC).isoformat(),
        "database": "postgresql",
        "run_id": run_id,
        "state": completed.state,
        "summary": expected_summary,
        "outputs": output_checks,
    }
    report_path = settings.data_dir / "migrations" / "stage10-pilot.json"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    print(
        json.dumps(
            {
                "verified": True,
                "run_id": run_id,
                "state": completed.state,
                "outputs": len(output_checks),
                "summary": expected_summary,
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
