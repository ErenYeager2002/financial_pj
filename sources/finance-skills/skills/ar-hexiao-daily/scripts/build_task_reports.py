#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把多个已选核销日的日期级报表整理为一份任务范围版静态 Excel。

批次内的日期级 Excel 只是逐日校验和写入时的临时产物。整合范围版成功
生成后，自动删除本次已选日期的日期级同类报告，只保留任务版交付文件。
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import common  # noqa: E402
import workbook_finalize  # noqa: E402


CLEANUP_REPORT_PREFIXES = ("核销日清", "变更清单", "订单写入差异")
AGGREGATED_SHEETS = {
    "核销明细": ("核销明细", "回款明细", "今日清单"),
    "流转表怎么填": ("流转表怎么填",),
}


def _dates(start: str, end: str):
    current = dt.date.fromisoformat(start)
    finish = dt.date.fromisoformat(end)
    if current > finish:
        raise ValueError("开始日期不能晚于结束日期")
    while current <= finish:
        yield current
        current += dt.timedelta(days=1)


def _selected_days(start: str, end: str, selected_dates=None) -> list[dt.date]:
    if selected_dates:
        days = sorted({dt.date.fromisoformat(value) for value in selected_dates})
        if start and days[0] < dt.date.fromisoformat(start):
            raise ValueError("已选日期不能早于开始日期")
        if end and days[-1] > dt.date.fromisoformat(end):
            raise ValueError("已选日期不能晚于结束日期")
        return days
    if not start or not end:
        raise ValueError("未提供已选日期时必须同时提供开始日期和结束日期")
    return list(_dates(start, end))


def _task_report_name(days: list[dt.date]) -> str:
    start = days[0].strftime("%Y%m%d")
    end = days[-1].strftime("%Y%m%d")
    consecutive = len(days) == (days[-1] - days[0]).days + 1
    if consecutive:
        return f"核销日清_{start}_{end}.xlsx"
    return f"核销日清_已选{len(days)}日_{start}_{end}.xlsx"


def _copy_sheet(source, target) -> None:
    for row in source.iter_rows():
        for cell in row:
            out = target[cell.coordinate]
            out.value = cell.value
            if isinstance(cell.value, str) and cell.value.startswith("="):
                out.data_type = "s"
            if cell.has_style:
                out._style = copy(cell._style)
            if cell.number_format:
                out.number_format = cell.number_format
            out.alignment = copy(cell.alignment)
            out.protection = copy(cell.protection)
    for key, dim in source.column_dimensions.items():
        target.column_dimensions[key].width = dim.width
        target.column_dimensions[key].hidden = dim.hidden
    for key, dim in source.row_dimensions.items():
        target.row_dimensions[key].height = dim.height
        target.row_dimensions[key].hidden = dim.hidden
    for merged in source.merged_cells.ranges:
        target.merge_cells(str(merged))
    target.freeze_panes = source.freeze_panes
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.auto_filter.ref = source.auto_filter.ref


def _daily_counts(out_dir: Path, day: dt.date) -> dict:
    token = day.strftime("%Y%m%d")
    final_path = out_dir / f"最终核销结果_{token}.json"
    if final_path.is_file():
        final = json.loads(final_path.read_text(encoding="utf-8"))
        if final.get("schema_version") != "ar-final-result-v1" or final.get("reconciliation_date") != day.isoformat():
            raise ValueError("最终核销结果的版本或日期无效")
        counts = final.get("counts") or {}
        flow = final.get("flow") or {}
        phases = flow.get("phases") or {}
        result = {"到账笔数": len({row.get("ar") for row in final.get("records", []) if row.get("ar")}),
                "订单行数": counts.get("total", 0),
                "自动": counts.get("completed", 0) + counts.get("skipped", 0),
                "挂账": counts.get("hold", 0), "异常": counts.get("exception", 0),
                "冲突": counts.get("conflict", 0), "实际写入记录数": final.get("write_count", 0),
                "写入 SO 数": counts.get("written_orders", 0), "最终跳过记录数": counts.get("skipped", 0),
                "统计口径": "最终处置；按首次判定记录计数", "首次校验跳过记录数": final.get("skip_count", ""),
                "流转人工填写数": _flow_count(flow.get("manual_count"), failed=True),
                "流转未完成日期数": int(any((phases.get(name) or {}).get("state") != "verified" for name in ("prefill", "status")))}
        for phase_name, label in (("prefill", "登记"), ("status", "状态")):
            phase = phases.get(phase_name) or {}
            for key, suffix in (("changed_count", "保留改动数"), ("unchanged_count", "无需改动数")):
                value = phase.get(key)
                result[f"流转{label}{suffix}"] = _flow_count(value, failed=phase.get("state") == "failed")
        return result
    result_path = out_dir / f"判定结果_{token}.json"
    if not result_path.is_file():
        return {}
    payload = json.loads(result_path.read_text(encoding="utf-8"))
    counts = payload.get("counts") or {}
    return {
        "到账笔数": payload.get("payment_count", ""),
        "订单行数": counts.get("total", ""),
        "自动": counts.get("auto", ""),
        "挂账": counts.get("hold", ""),
        "异常": counts.get("exception", ""),
        "统计口径": "旧版首次分类；未提供最终业务复核",
    }


def _flow_count(value, *, failed: bool) -> int | str:
    if isinstance(value, int) and not isinstance(value, bool) and value >= 0:
        return value
    return "读取失败" if failed else "未记录"


def _daily_report_files(out_dir: Path, prefix: str, day: dt.date) -> list[Path]:
    """返回某核销日的正式日报及多年度内部报告，不误匹配范围版文件。"""
    token = day.strftime("%Y%m%d")
    name_re = re.compile(
        rf"^{re.escape(prefix)}_{token}(?:_\d{{4}})?\.xlsx$"
    )
    return sorted(
        path for path in out_dir.glob(f"{prefix}_{token}*.xlsx")
        if name_re.fullmatch(path.name)
    )


def _append_aggregated_sheet(
    target,
    source,
    day: dt.date,
    header_written: bool,
) -> bool:
    rows = list(source.iter_rows())
    if not rows:
        return header_written
    source_header = rows[0]
    headings = [str(cell.value or "") for cell in source_header]
    if not header_written:
        target.append(["核销日期", *headings])
    else:
        existing = [cell.value for cell in target[1]]
        for heading in headings:
            if heading not in existing:
                existing.append(heading)
                target.cell(row=1, column=len(existing), value=heading)
    columns = {cell.value: cell.column for cell in target[1]}
    for cell in target[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for source_row in rows[1:]:
        if not any(cell.value not in (None, "") for cell in source_row):
            continue
        output_row = target.max_row + 1
        target.cell(row=output_row, column=1, value=day.isoformat())
        for heading, cell in zip(headings, source_row):
            output = target.cell(row=output_row, column=columns[heading], value=cell.value)
            if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                output.data_type = "s"
            if cell.has_style:
                output._style = copy(cell._style)
            output.number_format = cell.number_format
            output.alignment = copy(cell.alignment)
            output.protection = copy(cell.protection)
    return True


def build(
    workspace: Path,
    start: str,
    end: str,
    *,
    selected_dates=None,
    empty_dates=None,
) -> list[Path]:
    workspace = Path(workspace)
    if not (workspace / "04_产出").is_dir():
        workspace = common.resolve_workspace(workspace)
    out_dir = workspace / "04_产出"
    days = _selected_days(start, end, selected_dates)
    empty_days = {dt.date.fromisoformat(value) for value in (empty_dates or [])}
    unexpected_empty_days = empty_days - set(days)
    if unexpected_empty_days:
        raise ValueError("空核销日期必须属于本批已选日期")
    outputs = []
    daily_sources: set[Path] = set()
    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = "任务范围"
    headers = ["核销日期", "日报状态", "到账笔数", "订单行数", "自动", "挂账", "异常"]
    execution_columns = ["统计口径", "冲突", "实际写入记录数", "写入 SO 数", "首次校验跳过记录数", "最终跳过记录数",
                         "流转登记保留改动数", "流转登记无需改动数", "流转状态保留改动数", "流转状态无需改动数",
                         "流转人工填写数", "流转未完成日期数"] if empty_days or any(
        (out_dir / f"最终核销结果_{day.strftime('%Y%m%d')}.json").is_file() for day in days
    ) else []
    headers.extend(execution_columns)
    summary.append(headers)
    for cell in summary[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    summary.row_dimensions[1].height = 42

    aggregated = {title: wb.create_sheet(title) for title in AGGREGATED_SHEETS}
    header_written: set[str] = set()
    for day in days:
        token = day.strftime("%Y%m%d")
        daily = out_dir / f"核销日清_{token}.xlsx"
        for prefix in CLEANUP_REPORT_PREFIXES:
            daily_sources.update(_daily_report_files(out_dir, prefix, day))
        counts = _daily_counts(out_dir, day)
        if day in empty_days:
            counts = {column: 0 for column in [*headers[2:], *execution_columns]}
            counts["统计口径"] = "取数确认无记录；未执行核销写入"
        status = (
            "已纳入"
            if daily.is_file()
            else "无核销记录，已跳过"
            if day in empty_days
            else "当日无该报表"
        )
        summary.append([
            day.isoformat(), status,
            counts.get("到账笔数", ""), counts.get("订单行数", ""),
            counts.get("自动", ""), counts.get("挂账", ""), counts.get("异常", ""),
            *(counts.get(column, "") for column in execution_columns),
        ])
        if not daily.is_file():
            if day not in empty_days:
                raise FileNotFoundError(f"缺少日期级核销日清：{daily}")
            continue
        source_wb = openpyxl.load_workbook(daily, data_only=False, read_only=False)
        for target_title, aliases in AGGREGATED_SHEETS.items():
            source_title = next((name for name in aliases if name in source_wb.sheetnames), None)
            if source_title is None:
                continue
            target = aggregated[target_title]
            if _append_aggregated_sheet(target, source_wb[source_title], day, target_title in header_written):
                header_written.add(target_title)
        source_wb.close()

    summary.freeze_panes = "A2"
    summary.auto_filter.ref = summary.dimensions
    for col, width in zip("ABCDEFG", (14, 16, 12, 12, 10, 10, 10)):
        summary.column_dimensions[col].width = width
    for index, column in enumerate(execution_columns, start=8):
        summary.column_dimensions[openpyxl.utils.get_column_letter(index)].width = 38 if column == "统计口径" else 18
    for title, sheet in aggregated.items():
        if title not in header_written:
            sheet.append(["核销日期", "单号" if title == "核销明细" else "到账号(AR)", "状态"])
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.column_dimensions["A"].width = 14
    target = out_dir / _task_report_name(days)
    wb.save(target)
    workbook_finalize.finalize_static_report(target)
    outputs.append(target)

    # 必须等整合范围版保存并完成静态化后再删除日报，这样任一步失败时仍保留
    # 逐日证据，可安全重试。单日批次同样只交付这一份整合表。
    final_paths = {path.resolve() for path in outputs}
    for source in sorted(daily_sources):
        if source.resolve() not in final_paths and source.is_file():
            source.unlink()
    return outputs


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        description="生成多个已选核销日的一份整合核销日清，并移除已选日期的日报"
    )
    parser.add_argument("--workspace", default=str(common.WORK))
    parser.add_argument("--date-from", default="")
    parser.add_argument("--date-to", default="")
    parser.add_argument("--date", action="append", default=[], help="明确指定一个核销日，可重复")
    parser.add_argument(
        "--empty-date",
        action="append",
        default=[],
        help="平台已根据不可变取数包确认无核销记录的日期，可重复",
    )
    args = parser.parse_args(argv)
    try:
        outputs = build(
            Path(args.workspace),
            args.date_from,
            args.date_to,
            selected_dates=args.date,
            empty_dates=args.empty_date,
        )
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    for path in outputs:
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
