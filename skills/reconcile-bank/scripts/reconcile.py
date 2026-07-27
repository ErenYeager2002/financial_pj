from __future__ import annotations

import argparse
import json
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

DATE_ALIASES = ["交易日期", "凭证日期", "记账日期", "入账日期", "到账日期", "日期"]
AMOUNT_ALIASES = ["收入金额", "贷方发生额", "借方金额", "交易金额", "金额"]
REFERENCE_ALIASES = ["凭证号", "凭证编号", "流水号", "交易流水号", "摘要"]


@dataclass
class Record:
    row_number: int
    business_date: date
    amount: Decimal
    reference: str
    values: dict[str, Any]


def event(message: str, progress: int, *, state: str = "running") -> None:
    print(
        json.dumps(
            {"type": "progress", "message": message, "progress": progress, "state": state},
            ensure_ascii=False,
        ),
        flush=True,
    )


def normalized(value: Any) -> str:
    return str(value or "").strip().replace(" ", "").replace("\n", "")


def locate_header(headers: list[str], aliases: list[str]) -> int | None:
    normalized_headers = [normalized(value) for value in headers]
    for alias in aliases:
        if alias in normalized_headers:
            return normalized_headers.index(alias)
    for alias in aliases:
        for index, header in enumerate(normalized_headers):
            if len(alias) >= 2 and alias in header:
                return index
    return None


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


def parse_amount(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    text = str(value).replace(",", "").replace("￥", "").strip()
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return None


def read_records(path: Path, label: str) -> tuple[list[Record], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers_raw = next(rows, None)
    if not headers_raw:
        raise ValueError(f"{label}没有表头。")
    headers = [str(value or "").strip() for value in headers_raw]
    date_index = locate_header(headers, DATE_ALIASES)
    amount_index = locate_header(headers, AMOUNT_ALIASES)
    reference_index = locate_header(headers, REFERENCE_ALIASES)
    if date_index is None:
        raise ValueError(f"{label}没有找到日期列，支持：{DATE_ALIASES}")
    if amount_index is None:
        raise ValueError(f"{label}没有找到金额列，支持：{AMOUNT_ALIASES}")
    records: list[Record] = []
    warnings: list[str] = []
    for row_number, row in enumerate(rows, start=2):
        business_date = parse_date(row[date_index] if date_index < len(row) else None)
        amount = parse_amount(row[amount_index] if amount_index < len(row) else None)
        if business_date is None and amount is None:
            continue
        if business_date is None or amount is None:
            warnings.append(f"{label}第 {row_number} 行日期或金额无效，未参与匹配。")
            continue
        values = {
            headers[index] or f"列{index + 1}": value
            for index, value in enumerate(row)
            if index < len(headers)
        }
        reference = (
            str(row[reference_index] or "")
            if reference_index is not None and reference_index < len(row)
            else ""
        )
        records.append(
            Record(
                row_number=row_number,
                business_date=business_date,
                amount=amount,
                reference=reference,
                values=values,
            )
        )
    workbook.close()
    return records, warnings


def reconcile(
    bank_records: list[Record],
    ledger_records: list[Record],
    amount_tolerance: Decimal,
    date_tolerance_days: int,
) -> tuple[list[tuple[Record, Record, Decimal, int]], list[Record], list[Record], list[str]]:
    matched: list[tuple[Record, Record, Decimal, int]] = []
    used_ledger: set[int] = set()
    unmatched_bank: list[Record] = []
    warnings: list[str] = []
    for bank in bank_records:
        candidates: list[tuple[int, Decimal, int, int, Record]] = []
        for index, ledger in enumerate(ledger_records):
            if index in used_ledger:
                continue
            amount_diff = abs(bank.amount - ledger.amount)
            date_diff = abs((bank.business_date - ledger.business_date).days)
            if amount_diff <= amount_tolerance and date_diff <= date_tolerance_days:
                candidates.append((date_diff, amount_diff, ledger.row_number, index, ledger))
        if not candidates:
            unmatched_bank.append(bank)
            continue
        candidates.sort(key=lambda item: (item[0], item[1], item[2]))
        date_diff, amount_diff, _, ledger_index, ledger = candidates[0]
        used_ledger.add(ledger_index)
        matched.append((bank, ledger, amount_diff, date_diff))
        if len(candidates) > 1:
            warnings.append(
                f"银行流水第 {bank.row_number} 行有 {len(candidates)} 个候选，"
                "已按日期差、金额差和行号选择最优项。"
            )
    unmatched_ledger = [
        record for index, record in enumerate(ledger_records) if index not in used_ledger
    ]
    return matched, unmatched_bank, unmatched_ledger, warnings


def style_sheet(sheet: Any) -> None:
    fill = PatternFill("solid", fgColor="172033")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 32)
        sheet.column_dimensions[column[0].column_letter].width = max(width, 12)


def write_result(
    output_path: Path,
    matched: list[tuple[Record, Record, Decimal, int]],
    unmatched_bank: list[Record],
    unmatched_ledger: list[Record],
) -> None:
    workbook = Workbook()
    match_sheet = workbook.active
    match_sheet.title = "匹配明细"
    match_sheet.append(
        [
            "银行行号",
            "银行日期",
            "银行金额",
            "银行参考号",
            "总账行号",
            "总账日期",
            "总账金额",
            "凭证号",
            "金额差",
            "日期差（天）",
        ]
    )
    for bank, ledger, amount_diff, date_diff in matched:
        match_sheet.append(
            [
                bank.row_number,
                bank.business_date,
                float(bank.amount),
                bank.reference,
                ledger.row_number,
                ledger.business_date,
                float(ledger.amount),
                ledger.reference,
                float(amount_diff),
                date_diff,
            ]
        )
    bank_sheet = workbook.create_sheet("银行未匹配")
    bank_sheet.append(["原始行号", "日期", "金额", "参考号"])
    for item in unmatched_bank:
        bank_sheet.append([item.row_number, item.business_date, float(item.amount), item.reference])
    ledger_sheet = workbook.create_sheet("总账未匹配")
    ledger_sheet.append(["原始行号", "日期", "金额", "凭证号"])
    for item in unmatched_ledger:
        ledger_sheet.append(
            [item.row_number, item.business_date, float(item.amount), item.reference]
        )
    for sheet in workbook.worksheets:
        style_sheet(sheet)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--request", required=True)
    parser.add_argument("--result", required=True)
    args = parser.parse_args()

    request_path = Path(args.request).resolve()
    result_path = Path(args.result).resolve()
    payload = json.loads(request_path.read_text(encoding="utf-8"))
    parameters = payload["parameters"]
    amount_tolerance = Decimal(str(parameters["amount_tolerance"]))
    date_tolerance_days = int(parameters["date_tolerance_days"])
    bank_path = Path(payload["files"]["bank_file"]["local_path"])
    ledger_path = Path(payload["files"]["ledger_file"]["local_path"])

    event("正在读取银行流水", 15)
    bank_records, bank_warnings = read_records(bank_path, "银行流水")
    event("正在读取财务总账", 30)
    ledger_records, ledger_warnings = read_records(ledger_path, "财务总账")
    event("正在执行逐笔匹配", 55)
    matched, unmatched_bank, unmatched_ledger, match_warnings = reconcile(
        bank_records,
        ledger_records,
        amount_tolerance,
        date_tolerance_days,
    )
    event("正在生成对账结果", 80)
    output_path = Path(payload["output_dir"]) / "银行流水对账结果.xlsx"
    write_result(output_path, matched, unmatched_bank, unmatched_ledger)
    warnings = bank_warnings + ledger_warnings + match_warnings
    result = {
        "status": "success",
        "summary": {
            "bank_records": len(bank_records),
            "ledger_records": len(ledger_records),
            "matched": len(matched),
            "unmatched_bank": len(unmatched_bank),
            "unmatched_ledger": len(unmatched_ledger),
        },
        "output_files": [{"name": output_path.name, "path": str(output_path.resolve())}],
        "warnings": warnings,
    }
    result_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    event("对账结果已经生成", 95)


if __name__ == "__main__":
    main()
