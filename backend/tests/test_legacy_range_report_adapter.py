from pathlib import Path

import openpyxl

from app.legacy_range_report_adapter import build_selected_reports, main


def test_legacy_report_builder_aggregates_only_selected_sparse_dates(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    output_dir = workspace / "04_产出"
    output_dir.mkdir(parents=True)
    selected_dates = ["2026-08-07", "2026-08-10"]
    for selected_date in selected_dates:
        token = selected_date.replace("-", "")
        workbook = openpyxl.Workbook()
        workbook.active.title = "今日清单"
        workbook.active.append(["订单号", "金额"])
        workbook.active.append([f"SO-{token}", 100])
        workbook.save(output_dir / f"核销日清_{token}.xlsx")

    legacy_script = (
        Path(__file__).resolve().parents[2]
        / "skills"
        / "ar-hexiao-daily"
        / "vendor"
        / "scripts"
        / "build_task_reports.py"
    )

    report = build_selected_reports(legacy_script, workspace, selected_dates)

    assert report.name == "核销日清_已选2日_20260807_20260810.xlsx"
    workbook = openpyxl.load_workbook(report, data_only=True, read_only=True)
    try:
        summary = workbook["任务范围"]
        assert [summary.cell(row=row, column=1).value for row in (2, 3)] == selected_dates
        assert summary.cell(row=4, column=1).value is None
    finally:
        workbook.close()


def test_legacy_report_adapter_error_does_not_expose_paths(tmp_path: Path, capsys) -> None:
    secret_path = tmp_path / "private-workflow" / "build_task_reports.py"

    exit_code = main(
        [
            "--script",
            str(secret_path),
            "--workspace",
            str(tmp_path / "private-workspace"),
            "--date",
            "2026-08-04",
        ]
    )

    captured = capsys.readouterr()
    assert exit_code == 2
    assert captured.err == "ERROR: 旧版范围报告生成失败。\n"
    assert str(tmp_path) not in captured.err
