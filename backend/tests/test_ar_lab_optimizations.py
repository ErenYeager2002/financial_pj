"""Synthetic regression cases for the optional AR lab; never run live jobs."""
import datetime as dt
import importlib.util
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from app.ar_lab_execution import cached_command
from app.ar_skill_identity import AR_LAB_SKILL_ID


@pytest.fixture
def cache_module():
    path = Path(__file__).resolve().parents[2] / "skills/ar-hexiao-daily-lab/vendor/scripts/workbook_read_cache.py"
    spec = importlib.util.spec_from_file_location("isolated_ar_lab_cache", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def workbook(path, value):
    book = openpyxl.Workbook()
    book.active.title = "明细"
    book.active.append(["SO", "金额", "日期"])
    book.active.append(["SO-DEMO", value, dt.datetime(2026, 9, 4)])
    book.save(path)
    book.close()


def test_cache_preserves_types_and_does_not_return_mutable_shared_rows(tmp_path, cache_module, monkeypatch):
    path, cache = tmp_path / "ledger.xlsx", tmp_path / "rows.gz"
    workbook(path, 100)
    expected = cache_module.read_rows(path, "明细")
    cache_module.build(tmp_path, [(path, "明细")], cache)
    cache_module.configure(cache, cache_module.digest(cache), tmp_path)
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: pytest.fail("cache should avoid XML parsing"))
    first = cache_module.read_rows(path, "明细")
    assert first == expected
    first.clear()
    assert cache_module.read_rows(path, "明细") == expected


def test_changed_workbook_never_uses_prewrite_cache(tmp_path, cache_module):
    path, cache = tmp_path / "ledger.xlsx", tmp_path / "rows.gz"
    workbook(path, 100)
    cache_module.build(tmp_path, [(path, "明细")], cache)
    cache_module.configure(cache, cache_module.digest(cache), tmp_path)
    workbook(path, 200)
    assert cache_module.read_rows(path, "明细")[1][1] == 200


def test_changed_cache_and_other_task_root_are_rejected(tmp_path, cache_module):
    path, cache = tmp_path / "ledger.xlsx", tmp_path / "rows.gz"
    workbook(path, 100)
    cache_module.build(tmp_path, [(path, "明细")], cache)
    fingerprint = cache_module.digest(cache)
    with pytest.raises(ValueError):
        cache_module.configure(cache, fingerprint, tmp_path / "another-task")
    cache.write_bytes(cache.read_bytes() + b"changed")
    with pytest.raises(ValueError):
        cache_module.configure(cache, fingerprint, tmp_path)


def test_uncached_flow_sheet_does_not_hash_whole_workbook(tmp_path, cache_module, monkeypatch):
    cache_module._root = tmp_path
    cache_module._entries = {"digest:明细": []}
    monkeypatch.setattr(cache_module, "digest", lambda *a: pytest.fail("unrelated sheet has no cache candidate"))
    worksheet = SimpleNamespace(iter_rows=lambda **kwargs: iter([("flow",)]))
    assert cache_module.read_rows(tmp_path / "flow.xlsx", "9月", worksheet=worksheet) == [("flow",)]


def test_later_batch_date_uses_verified_shared_storage_root(tmp_path):
    shared = tmp_path / "first-date"
    shared.mkdir()
    execution = SimpleNamespace(
        workflow=SimpleNamespace(skill_id=AR_LAB_SKILL_ID, id="later-date"), db=None,
        context={"ar_read_cache": {"path": str(shared / "read-cache" / "later-action.gz"), "fingerprint": "fixed"}},
        service=SimpleNamespace(_workflow_storage_root=lambda db, workflow: shared),
    )
    name, arguments = cached_command(execution, "classify_hexiao.py", ["--workspace", str(shared / "workspace")])
    assert name == "run_read_cached.py" and arguments[1] == str(shared)
    assert cached_command(execution, "verify_execution_write.py", ["--workspace", str(shared)])[0] == "verify_execution_write.py"
    execution.workflow.skill_id = "ar-hexiao-daily"
    assert cached_command(execution, "classify_hexiao.py", ["original"]) == ("classify_hexiao.py", ["original"])
