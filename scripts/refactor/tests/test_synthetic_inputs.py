import os
from pathlib import Path
import sys
import tempfile
import unittest
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from synthetic_inputs import create
from openpyxl import load_workbook

class SyntheticInputTests(unittest.TestCase):
    def test_generated_workbooks_roundtrip_and_current_export_reader(self):
        with tempfile.TemporaryDirectory(prefix='financial-refactor-fixture-') as temp:
            base=Path(temp);(base/'.refactor-isolated').write_text('synthetic-only-v1')
            env={'REFACTOR_TEST_ROOT':str(base),'FINANCIAL_ENV':'test','FINANCIAL_DATA_DIR':str(base/'data'),'FINANCIAL_DATABASE_URL':'sqlite:///'+str(base/'test.db'),'FINANCIAL_AR_HEXIAO_EXECUTION_ENABLED':'false','FINANCIAL_TASK_DISCOVERY_ENABLED':'false','REFACTOR_REAL_CONNECTORS':'disabled'}
            root=base/'inputs';manifest=create(root,env)
            self.assertEqual(len(manifest['sha256']),11)
            for year,name in manifest['ledger_years'].items():
                wb=load_workbook(root/name);ws=wb['明细']
                self.assertEqual(ws['F2'].value,100)
                self.assertIsNone(ws['H2'].value)
                self.assertIn(year[2:],ws['E2'].value);wb.close()
            scripts=Path(__file__).resolve().parents[3]/'skills/ar-hexiao-daily/vendor/scripts'
            sys.path.insert(0,str(scripts))
            try:
                import classify_hexiao
                for day in manifest['dates']:
                    payments=classify_hexiao.load_exports(root/day)
                    self.assertEqual(len(payments),1)
                    self.assertEqual(sum(payments[0]['writeoffs'].values()),100)
            finally: sys.path.remove(str(scripts))
            with self.assertRaises(FileExistsError): create(root,env)
