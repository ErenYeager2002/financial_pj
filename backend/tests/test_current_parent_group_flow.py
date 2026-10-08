import copy
import datetime as dt
import tempfile
import unittest
from pathlib import Path
import openpyxl
import build_flow_plan as B
import validate_plan as V
import apply_flow
from current_parent_receipt_group import prove,check
from test_current_parent_receipt_group import ParentReceiptGroupTest

class ParentGroupFlowTest(unittest.TestCase):
    def setup_case(self,root):
        books=root/'02_我的表副本';books.mkdir();(root/'04_产出').mkdir()
        parents,_=ParentReceiptGroupTest().fixture()
        so='SO26010001'
        for p in parents:p['orders'][0]['so']=so
        book=books/'2026年盈亏.xlsx'
        wb=openpyxl.Workbook();ws=wb.active;ws.title='明细'
        ws.append(['新智云单号','实收金额','应收金额','计提','回款明细','是否结账','收款时间','收款方式(支/汇/现)','差异'])
        ws.append([so,'SOD_A',900,900,900,'是',dt.date(2026,8,12),'汇',None]);wb.save(book);wb.close()
        rows=V.read_ledger_rows(book);proof=prove(parents,rows,'2026-08-25');assert proof
        flow=books/'flow.xlsx';wb=openpyxl.Workbook();ws=wb.active;ws.title='流水'
        ws.append(['日期','公司名称','金额','收款形式','单号','预收','是否更新应收款'])
        for _ in range(2):ws.append([dt.date(2026,8,12),'客户',450,'汇款',so,None,'是'])
        wb.save(flow);wb.close()
        members=[]
        for parent in parents:
            ar=parent['ar']
            r=dict(ar=ar,so=so,sod='SOD_A',case_id=ar+'|'+so,code='OK_SO_ALREADY_SETTLED',bucket='auto',
                _check={'verdict':'skip'},flow_hits=2,flow_matched_by='三键',ledger_path=str(book),
                write_currency_audit={'currency':'CNY','amount_local':0,'amount_orig':0},
                current_parent_receipt_group=proof,source_lineage={'source':dict(ar=ar,so=so,currency='CNY',
                    reconciliation_date='2026-08-25',parent_amount_orig=450,parent_amount_local=450)})
            r['flow_parent_group_proof']=check(proof,rows,r,'2026-08-25');assert r['flow_parent_group_proof']
            members.append(r)
        return book,flow,dict(hexiao_date='2026-08-25',write=[],skip=members,conflict=[])

    def plan(self,root,checked):
        return B.finalize_plan_after_ledger(B.build_plan({'auto':checked['skip'],'hexiao_date':'2026-08-25'}),checked,workspace=root)

    def test_physical_binding_repeat_reorder_and_preserve_pl(self):
        for reverse in (False,True):
            with self.subTest(reverse=reverse),tempfile.TemporaryDirectory() as temp:
                root=Path(temp);book,flow,checked=self.setup_case(root)
                if reverse:checked['skip'].reverse()
                pl=book.read_bytes();saved=copy.deepcopy(checked)
                final=self.plan(root,checked)
                self.assertTrue(all(i['verdict']=='write' for i in final['items']),final)
                changes,errors=apply_flow.write_flow_items(root,final['items'],in_place=True,phase='status')
                self.assertFalse(errors);self.assertEqual(len(changes),2)
                self.assertTrue(all(c['预收公式']=='450-450=0' for c in changes))
                self.assertEqual(saved,checked);self.assertEqual(book.read_bytes(),pl)
                raw=flow.read_bytes();checked['skip'].reverse()
                repeated=self.plan(root,checked)
                self.assertEqual(apply_flow.write_flow_items(root,repeated['items'],in_place=True,phase='status'),([],[]))
                self.assertEqual(flow.read_bytes(),raw)

    def test_unsafe_group_or_changed_evidence_never_gets_write_binding(self):
        for mode in ('missing_member','extra_row','different_prepay','changed_accrual','changed_amount','proof_tampered'):
            with self.subTest(mode=mode),tempfile.TemporaryDirectory() as temp:
                root=Path(temp);book,flow,checked=self.setup_case(root)
                if mode=='missing_member':checked['skip'].pop()
                if mode=='proof_tampered':checked['skip'][0]['flow_parent_group_proof']['amount']=900
                if mode in ('extra_row','different_prepay'):
                    wb=openpyxl.load_workbook(flow);ws=wb.active
                    if mode=='extra_row':ws.append([c.value for c in ws[2]])
                    else:ws['F3']=1
                    wb.save(flow);wb.close()
                if mode in ('changed_accrual','changed_amount'):
                    wb=openpyxl.load_workbook(book);wb.active['D2' if mode=='changed_accrual' else 'E2']=899;wb.save(book);wb.close()
                before=(book.read_bytes(),flow.read_bytes())
                final=self.plan(root,checked)
                self.assertTrue(all(i['verdict']!='write' for i in final['items']),final)
                self.assertEqual(before,(book.read_bytes(),flow.read_bytes()))

    def test_full_flow_review_rebuilds_from_fixed_annual_material(self):
        import json,shutil
        import execution_flow_stage as X
        import verify_execution_write as E
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);book,flow,checked=self.setup_case(root)
            baseline=root/'baseline';shutil.copytree(root/'02_我的表副本',baseline/'02_我的表副本')
            output=root/'04_产出'
            plan=B.build_plan({'auto':checked['skip'],'hexiao_date':'2026-08-25'})
            (output/'流转写入计划_校验后.json').write_text(json.dumps(plan,default=str))
            checked_path=output/'checked.json';checked_path.write_text(json.dumps(checked,default=str))
            result=X.run(root,checked_path)
            self.assertTrue(result['flow_written'],result)
            (output/'流转阶段执行结果.json').write_text(json.dumps(result,default=str))
            self.assertEqual(E.verify_flow(root,baseline,flow,checked)['mode'],'approved_patch_all_parts')
            actual=flow.read_bytes();original=book.read_bytes()
            baseline_book=baseline/'02_我的表副本'/book.name
            wb=openpyxl.load_workbook(baseline_book);wb.active['D2']=899;wb.save(baseline_book);wb.close()
            with self.assertRaises(ValueError):E.verify_flow(root,baseline,flow,checked)
            self.assertEqual(flow.read_bytes(),actual);self.assertEqual(book.read_bytes(),original)
