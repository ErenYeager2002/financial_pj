from __future__ import annotations

import argparse
import hashlib
import io
import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.database import SessionLocal
from app.models import AuditEvent, FileRecord, RunRecord, StepRun

EXPECTED_COUNTS = {
    "compliance-spot-check": (2, 2),
    "dept-expense-alloc": (4, 1),
    "dreame-ar-progress-diff": (2, 1),
    "labor-invoice-check": (2, 1),
    "order-daily-summary": (1, 1),
    "project-detail-to-ledger": (2, 2),
    "receivables-merge": (1, 1),
    "reconcile-bank": (2, 1),
    "split-by-sales": (1, 1),
}

EXPECTED_SHEETS = {
    "dept-expense-alloc": {"部门科目余额表", "利润表"},
    "project-detail-to-ledger": {"明细"},
    "receivables-merge": {"主表", "认列告警", "运行报告"},
    "reconcile-bank": {"匹配明细", "银行未匹配", "总账未匹配"},
}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def file_bindings(value: object) -> list[dict[str, object]]:
    if isinstance(value, dict):
        return [value]
    if isinstance(value, list):
        return [item for item in value if isinstance(item, dict)]
    return []


def inspect_xlsx(content: bytes) -> set[str]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    sheets = set(workbook.sheetnames)
    workbook.close()
    assert sheets
    return sheets


def inspect_output(skill_id: str, path: Path) -> dict[str, int]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        sheets = inspect_xlsx(path.read_bytes())
        required = EXPECTED_SHEETS.get(skill_id, set())
        assert required <= sheets, (skill_id, required - sheets)
        return {"xlsx": 1, "zip_entries": 0, "nested_xlsx": 0}
    if suffix == ".zip":
        with zipfile.ZipFile(path) as archive:
            assert archive.testzip() is None
            names = [name for name in archive.namelist() if not name.endswith("/")]
            assert names
            nested_xlsx = 0
            for name in names:
                if name.lower().endswith(".xlsx"):
                    inspect_xlsx(archive.read(name))
                    nested_xlsx += 1
            assert nested_xlsx >= 1
            return {"xlsx": 0, "zip_entries": len(names), "nested_xlsx": nested_xlsx}
    if suffix == ".json":
        assert isinstance(json.loads(path.read_text(encoding="utf-8")), dict)
    else:
        assert path.stat().st_size > 0
    return {"xlsx": 0, "zip_entries": 0, "nested_xlsx": 0}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run", action="append", required=True, metavar="SKILL_ID=RUN_ID")
    args = parser.parse_args()
    pairs = [item.split("=", 1) for item in args.run]
    assert {skill_id for skill_id, _ in pairs} == set(EXPECTED_COUNTS)

    results: list[dict[str, object]] = []
    with SessionLocal() as db:
        for skill_id, run_id in pairs:
            run = db.get(RunRecord, run_id)
            assert run is not None
            assert run.skill_id == skill_id
            assert run.state == "succeeded"
            assert run.progress == 100
            assert run.confirmed_at is not None

            inputs = [
                binding
                for value in json.loads(run.files_json).values()
                for binding in file_bindings(value)
            ]
            expected_inputs, expected_outputs = EXPECTED_COUNTS[skill_id]
            assert len(inputs) == expected_inputs
            input_paths: set[Path] = set()
            for binding in inputs:
                record = db.get(FileRecord, str(binding["file_id"]))
                assert record is not None and record.kind == "input"
                path = Path(record.stored_path).resolve()
                assert path.is_file()
                assert digest(path) == record.sha256 == binding["sha256"]
                input_paths.add(path)

            output_items = json.loads(run.result_json)["output_files"]
            assert len(output_items) == expected_outputs
            output_stats = {"xlsx": 0, "zip_entries": 0, "nested_xlsx": 0}
            for item in output_items:
                record = db.get(FileRecord, item["file_id"])
                assert record is not None and record.kind == "output"
                assert record.run_id == run_id
                path = Path(record.stored_path).resolve()
                assert path.is_file() and path not in input_paths
                assert digest(path) == record.sha256 == item["sha256"]
                inspected = inspect_output(skill_id, path)
                for key, value in inspected.items():
                    output_stats[key] += value
                downloads = db.scalar(
                    select(func.count(AuditEvent.id)).where(
                        AuditEvent.action == "file.download",
                        AuditEvent.resource_id == record.id,
                        AuditEvent.outcome == "success",
                    )
                )
                assert int(downloads or 0) >= 1

            steps = db.scalars(select(StepRun).where(StepRun.run_id == run_id)).all()
            assert len(steps) == 5
            assert all(step.state == "succeeded" for step in steps)
            results.append(
                {
                    "skill_id": skill_id,
                    "run_id": run_id,
                    "version": run.skill_version,
                    "inputs": len(inputs),
                    "outputs": len(output_items),
                    "steps": len(steps),
                    **output_stats,
                }
            )
    print(json.dumps(results, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
