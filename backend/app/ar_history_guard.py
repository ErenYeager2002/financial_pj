"""Apply material differences to individual decisions, including pinned old tasks."""
import argparse
import json
from pathlib import Path
import sys


def guard_decisions(result, differences, classifier):
    conflicts = {}
    for item in differences:
        if item['status'] == 'conflict':
            conflicts.setdefault((item['year'], item['so']), []).append(item)
    records = [r for bucket in ('auto', 'hold', 'exception') for r in result.get(bucket, [])]
    for row in records:
        differences_for_order = conflicts.get((row.get('ledger_year', row.get('target_ledger_year')), row.get('so')), [])
        if not differences_for_order:
            continue
        # Prior reports describe prior materials. They are evidence to review,
        # never an independent veto over a plan derived from the current file.
        row['material_history_conflicts'] = differences_for_order
    if hasattr(classifier, 'FS'):
        classifier.FS.guard(records)
    for bucket in ('auto', 'hold', 'exception'):
        result[bucket] = [r for r in records if r['bucket'] == bucket]
    result['counts'] = {bucket: len(result[bucket]) for bucket in ('auto', 'hold', 'exception')}
    result['counts']['total'] = len(records)
    result['e_code_dist'] = classifier._dist(records)
    result['ar_summary'] = classifier.build_ar_summary(records)
    result['material_history_differences'] = differences


def append_history_report(workspace, tag, differences):
    import openpyxl
    from openpyxl.styles import Alignment, Font
    result_path = workspace / '04_产出' / f'最终核销结果_{tag}.json'
    payload = json.loads(result_path.read_text(encoding='utf-8'))
    written = {r['case_id'] for r in payload.get('records', []) if r.get('execution_status') == 'written_verified'}
    checked_path = workspace / '04_产出' / f'写入计划_校验后_{tag}.json'
    checked = json.loads(checked_path.read_text(encoding='utf-8')) if checked_path.is_file() else {}
    corrected = {(r.get('ledger_year'), r.get('so'), r.get('sod')) for r in checked.get('write', [])
                 if r.get('receipt_correction') and r.get('case_id') in written}
    path = workspace / '04_产出' / f'核销日清_{tag}.xlsx'
    book = openpyxl.load_workbook(path)
    if '历史材料差异' in book.sheetnames:
        del book['历史材料差异']
    sheet = book.create_sheet('历史材料差异')
    sheet.append(['年度', '单号', 'SOD', '处理', '原记录行', '当前行', '历史金额', '历史日期', '历史方式', '当前记录', '说明'])
    for d in differences:
        old = d['expected']
        sheet.append([d['year'], d['so'], d['sod'], '已按本次核销结果覆盖并回读' if (d['year'],d['so'],d['sod']) in corrected else ('历史值与当前表不同，以本次核销判定为准' if d['status'] == 'conflict' else '缺少历史记录，按本次取数逐条判定'),
                      ','.join(map(str, d['previous_rows'])), ','.join(str(r['row']) for r in d['current_rows']),
                      old['amount'], old['date'], old['method'],
                      '\n'.join(f"第{r['row']}行：{r['amount']} / {r['date']} / {r['method']}" for r in d['current_rows']),
                      '历史差异清单不代表本次已补写；实际写入、跳过或挂账以本次核销明细为准'])
    sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]: cell.font = Font(bold=True)
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.data_type == 'f': cell.data_type = 's'
    for col in ('A','B','C','D','E','F','G','H','I','J','K'):
        sheet.column_dimensions[col].width = 24 if col not in ('D','J','K') else 48
    book.save(path); book.close()
    result_path = workspace / '04_产出' / f'最终核销结果_{tag}.json'
    payload = json.loads(result_path.read_text(encoding='utf-8'))
    payload['material_history_differences'] = differences
    result_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')


def main():
    parser = argparse.ArgumentParser()
    for name in ('scripts', 'workspace', 'date', 'mode'): parser.add_argument('--' + name, required=True)
    args = parser.parse_args(); workspace = Path(args.workspace)
    history = workspace / '03_台账' / '业务材料历史差异.json'
    if not history.is_file(): return
    payload = json.loads(history.read_text(encoding='utf-8'))
    if payload.get('schema_version') != 1 or not isinstance(payload.get('differences'), list):
        raise ValueError('材料历史差异清单无效')
    differences = payload['differences']
    if not differences: return
    tag = args.date.replace('-', '')
    if args.mode == 'classify':
        sys.path.insert(0, args.scripts)
        import classify_hexiao as classifier
        path = workspace / '04_产出' / f'判定结果_{tag}.json'
        result = json.loads(path.read_text(encoding='utf-8'))
        guard_decisions(result, differences, classifier)
        path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    elif args.mode == 'report':
        append_history_report(workspace, tag, differences)
    else: raise ValueError('未知历史差异处理阶段')


if __name__ == '__main__': main()
