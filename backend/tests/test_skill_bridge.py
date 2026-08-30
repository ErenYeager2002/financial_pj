from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[2]


def test_legacy_bridge_executes_cli_and_collects_output(tmp_path: Path) -> None:
    skill_dir = tmp_path / "sample-skill"
    scripts_dir = skill_dir / "scripts"
    vendor_dir = skill_dir / "vendor" / "scripts"
    scripts_dir.mkdir(parents=True)
    vendor_dir.mkdir(parents=True)
    shutil.copy2(PROJECT_ROOT / "scripts" / "legacy_skill_bridge.py", scripts_dir / "entry.py")
    (vendor_dir / "fake.py").write_text(
        """
import argparse
from pathlib import Path

parser = argparse.ArgumentParser()
parser.add_argument("--source", required=True)
parser.add_argument("--note", required=True)
parser.add_argument("--out", required=True)
args = parser.parse_args()
Path(args.out).write_text(f"{Path(args.source).name}|{args.note}", encoding="utf-8")
""".strip(),
        encoding="utf-8",
    )
    (skill_dir / "bridge.yaml").write_text(
        yaml.safe_dump(
            {
                "name": "桥接测试",
                "command": "vendor/scripts/fake.py",
                "arguments": [
                    {"kind": "file", "role": "source", "flag": "--source"},
                    {"kind": "parameter", "name": "note", "flag": "--note"},
                ],
                "output": {"type": "file", "flag": "--out", "path": "结果.txt"},
            },
            allow_unicode=True,
            sort_keys=False,
        ),
        encoding="utf-8",
    )
    input_path = tmp_path / "原始输入.xlsx"
    input_path.write_bytes(b"test")
    output_dir = tmp_path / "run" / "outputs"
    request_path = tmp_path / "request.json"
    result_path = tmp_path / "result.json"
    request_path.write_text(
        json.dumps(
            {
                "parameters": {"note": "已处理"},
                "files": {"source": {"local_path": str(input_path)}},
                "output_dir": str(output_dir),
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    completed = subprocess.run(
        [
            sys.executable,
            str(scripts_dir / "entry.py"),
            "--request",
            str(request_path),
            "--result",
            str(result_path),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    result = json.loads(result_path.read_text(encoding="utf-8"))
    assert result["status"] == "success"
    assert result["summary"]["output_count"] == 1
    assert result["output_files"][0]["name"] == "结果.txt"
    assert (output_dir / "结果.txt").read_text(encoding="utf-8") == "原始输入.xlsx|已处理"


def test_withholding_report_rename_fails_without_recognized_pdf(tmp_path: Path) -> None:
    skill_dir = PROJECT_ROOT / "skills" / "withholding-report-rename"
    input_dir = tmp_path / "输入"
    input_dir.mkdir()
    (input_dir / "sample.pdf").write_bytes(b"not-a-real-pdf")
    output_dir = tmp_path / "run" / "outputs"
    request_path = tmp_path / "request.json"
    result_path = tmp_path / "result.json"
    request_path.write_text(
        json.dumps(
            {
                "files": {"reports": {"local_path": str(input_dir / "sample.pdf")}},
                "output_dir": str(output_dir),
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    completed = subprocess.run(
        [
            sys.executable,
            str(skill_dir / "scripts" / "entry.py"),
            "--request",
            str(request_path),
            "--result",
            str(result_path),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )

    assert completed.returncode != 0
    assert not result_path.exists()
    assert not list(output_dir.rglob("*.zip"))
    assert not list(output_dir.rglob("*.csv"))


def test_project_detail_bridge_appends_to_a_copy_and_preserves_inputs(tmp_path: Path) -> None:
    skill_dir = PROJECT_ROOT / "skills" / "project-detail-to-ledger"
    project_path = tmp_path / "项目明细.xlsx"
    ledger_path = tmp_path / "盈亏核算表.xlsx"

    project = Workbook()
    project_sheet = project.active
    project_sheet.title = "项目明细"
    project_sheet.append(
        [
            "销售",
            "客户",
            "SO",
            "SOD",
            "业务类别",
            "订单名称",
            "下单日期",
            "整单交付日期",
            "下单数量",
            "单价",
            "交付额/本币",
            "项目经理",
        ]
    )
    project_sheet.append(
        [
            "张三",
            "客户A",
            "SO-NEW",
            "SOD-NEW",
            "笔译",
            "项目A",
            "2026-08-01",
            "2026-08-02",
            100,
            2.5,
            250,
            "经理A",
        ]
    )
    project.save(project_path)

    ledger = Workbook()
    ledger_sheet = ledger.active
    ledger_sheet.title = "明细"
    ledger_sheet.append(
        [
            "销售人员",
            "客户名称",
            "新智云单号",
            "翻译类型",
            "文件名",
            "项目下单日期",
            "项目交付日期",
            "字数统计",
            "价格",
            "应收金额",
            "实收金额",
            "项目经理",
        ]
    )
    ledger_sheet.append(
        [
            "李四",
            "客户旧",
            "SO-OLD",
            "笔译",
            "旧项目",
            "2026-07-01",
            "2026-07-02",
            50,
            2,
            100,
            "SOD-OLD",
            "经理B",
        ]
    )
    ledger.save(ledger_path)

    project_before = project_path.read_bytes()
    ledger_before = ledger_path.read_bytes()
    output_dir = tmp_path / "run" / "outputs"
    request_path = tmp_path / "request.json"
    result_path = tmp_path / "result.json"
    request_path.write_text(
        json.dumps(
            {
                "files": {
                    "project_detail": {"local_path": str(project_path)},
                    "ledger": {"local_path": str(ledger_path)},
                },
                "parameters": {},
                "output_dir": str(output_dir),
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    completed = subprocess.run(
        [
            sys.executable,
            str(skill_dir / "scripts" / "entry.py"),
            "--request",
            str(request_path),
            "--result",
            str(result_path),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    assert project_path.read_bytes() == project_before
    assert ledger_path.read_bytes() == ledger_before
    result = json.loads(result_path.read_text(encoding="utf-8"))
    assert result["status"] == "success"
    assert result["summary"]["output_count"] == 2
    output_path = output_dir / "项目明细补录结果.xlsx"
    assert output_path.is_file()
    assert (output_dir / "项目明细补录结果_补录报告.json").is_file()
    output = load_workbook(output_path, data_only=False)
    assert output["明细"].cell(row=3, column=5).value == "SO-NEW"
    assert output["明细"].cell(row=3, column=20).value == "SOD-NEW"
    output.close()


def test_legacy_bridge_rejects_missing_required_outputs_before_archive(tmp_path: Path) -> None:
    skill_dir = tmp_path / "rename-skill"
    scripts_dir = skill_dir / "scripts"
    vendor_dir = skill_dir / "vendor" / "scripts"
    scripts_dir.mkdir(parents=True)
    vendor_dir.mkdir(parents=True)
    shutil.copy2(PROJECT_ROOT / "scripts" / "legacy_skill_bridge.py", scripts_dir / "entry.py")
    (vendor_dir / "fake.py").write_text(
        """
import argparse
from pathlib import Path

parser = argparse.ArgumentParser()
parser.add_argument("--input", required=True)
parser.add_argument("--mode", required=True)
parser.add_argument("--out-dir", required=True)
args = parser.parse_args()
Path(args.out_dir).mkdir(parents=True, exist_ok=True)
Path(args.out_dir, "对照表.csv").write_text("原文件名,新文件名\\n", encoding="utf-8")
""".strip(),
        encoding="utf-8",
    )
    (skill_dir / "bridge.yaml").write_text(
        yaml.safe_dump(
            {
                "name": "申报表重命名测试",
                "command": "vendor/scripts/fake.py",
                "arguments": [
                    {"kind": "input_dir", "role": "reports", "flag": "--input"},
                    {"kind": "static", "values": ["--mode", "copy"]},
                ],
                "output": {
                    "type": "directory",
                    "flag": "--out-dir",
                    "path": "重命名结果",
                    "required_globs": ["*.pdf"],
                    "archive": True,
                    "archive_name": "申报表结果.zip",
                },
            },
            allow_unicode=True,
            sort_keys=False,
        ),
        encoding="utf-8",
    )
    input_dir = tmp_path / "输入"
    input_dir.mkdir()
    input_path = input_dir / "sample.pdf"
    input_path.write_bytes(b"test")
    output_dir = tmp_path / "run" / "outputs"
    request_path = tmp_path / "request.json"
    result_path = tmp_path / "result.json"
    request_path.write_text(
        json.dumps(
            {
                "files": {"reports": {"local_path": str(input_path)}},
                "output_dir": str(output_dir),
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    completed = subprocess.run(
        [
            sys.executable,
            str(scripts_dir / "entry.py"),
            "--request",
            str(request_path),
            "--result",
            str(result_path),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )

    assert completed.returncode != 0
    assert not result_path.exists()
    assert not (output_dir / "申报表结果.zip").exists()
