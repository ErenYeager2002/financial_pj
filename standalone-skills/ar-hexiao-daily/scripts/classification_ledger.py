"""Read-only workbook indexes and current-row matching."""
from __future__ import annotations

from pathlib import Path
from typing import Dict
from typing import List
from typing import Optional
from typing import Tuple
import baseline_receipts as BR
import common
import datetime as dt
import settlement_status
from classification_contract import TOL


class LedgerIndex:
    """盈亏表『明细』只读索引。她表里**一行 = 一个 SOD**。"""

    def __init__(self, path: Optional[Path] = None, synthetic: Optional[dict] = None):
        self.path = path
        self.so_index: Dict[str, List[int]] = {}
        self.sod_index: Dict[str, List[int]] = {}
        self.so_amount_index: Dict[Tuple[str, int], List[int]] = {}
        self.row_snapshot: Dict[int, dict] = {}
        self.cols: dict = {}
        if synthetic is not None:
            self.so_index = {k: list(v) for k, v in synthetic.get("so", {}).items()}
            self.sod_index = {k: list(v) for k, v in synthetic.get("sod", {}).items()}
            self.row_snapshot = synthetic.get("rows", {})
            for r, snap in self.row_snapshot.items():
                y = common.to_number(snap.get("yingshou"))
                so = snap.get("so") or ""
                if so and y is not None:
                    self.so_amount_index.setdefault((so, int(round(y * 100))), []).append(int(r))
            return
        if path is not None:
            self._load(path)

    def _load(self, path: Path):
        import openpyxl

        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        if "明细" not in wb.sheetnames:
            names = list(wb.sheetnames)
            wb.close()
            raise ValueError(f"盈亏表无『明细』sheet：{names}")
        all_rows = list(wb["明细"].iter_rows(values_only=True))
        wb.close()

        aliases = common.load_aliases()
        hrow, headers = common.find_header_row(
            all_rows, "盈亏明细", ["SO", "SOD", "计提", "回款明细", "是否结账"], aliases
        )
        cols = common.resolve_columns(
            headers, "盈亏明细",
            ["SO", "SOD", "计提", "回款明细", "是否结账", "收款时间", "收款方式"],
            aliases,
        )
        role = aliases.get("盈亏明细", {})
        diff_idx = common.fuzzy_find_col(headers, role.get("差异", ["差异"]))
        if diff_idx is not None:
            cols["差异"] = diff_idx
        yidx = common.fuzzy_find_col(headers, role.get("应收", ["应收金额", "应收"]))
        if yidx is not None:
            cols["应收"] = yidx
        self.cols = cols

        for r_i, row in enumerate(all_rows, start=1):
            if r_i <= hrow + 1:
                continue
            vals = list(row)

            def cell(key):
                i = cols.get(key)
                return vals[i] if i is not None and i < len(vals) else None

            so_s = str(cell("SO") or "").strip()
            sod_s = str(cell("SOD") or "").strip()
            if not so_s.startswith("SO") and not sod_s.startswith("SOD"):
                continue
            if so_s.startswith("SO"):
                self.so_index.setdefault(so_s, []).append(r_i)
            if sod_s.startswith("SOD"):
                self.sod_index.setdefault(sod_s, []).append(r_i)
            yingshou = common.to_number(cell("应收"))
            if so_s.startswith("SO") and yingshou is not None:
                self.so_amount_index.setdefault(
                    (so_s, int(round(yingshou * 100))), []
                ).append(r_i)
            self.row_snapshot[r_i] = {
                "so": so_s, "sod": sod_s,
                "jiti": cell("计提"), "huikuan": cell("回款明细"),
                "chayi": cell("差异"),
                "jiezhang": cell("是否结账"), "shoukuan_time": cell("收款时间"),
                "shoukuan_way": cell("收款方式"), "yingshou": yingshou,
            }

    def positional_row(
        self, so: str, sod: str, lines: List[dict]
    ) -> Optional[Tuple[int, str, Optional[float]]]:
        """
        (SO+应收金额) 定位不到唯一行时的严格消歧：验证**整段能不能对齐**。

          她表里这个 SO 的全部行（行号升序） vs 智云该 SO 的全部 SOD（编号降序）

        为什么按这个序：两边都源自月初同一份「交付数据」导出，天然同序。

        返回 (行号, 依据, 比例)；对不齐返回 None（老老实实挂起）。依据两种：

        · `exact`  —— 金额序列逐位完全相等。实测 SO26040297（12 行）、SO26040481（10 行）。
        · `ratio`  —— 有几位不等，但**所有不等的位比值完全一致**（系统性口径差，不是行错位）。
                      实测 SO26040322：智云 477.61/661.31 vs 她表 488.64/676.58，
                      两处比值都是 0.977433，她当天就是按智云金额填的。
                      行错位不可能凑出同一个比值，所以这条判据是可证的，不是猜。

        只要行数对不上（比如她把某个 SOD 拆成了两行）→ 直接 None。
        """
        rows = sorted(self.so_index.get(so, []))
        if not rows or not lines or len(rows) != len(lines):
            return None
        ordered = sorted(lines, key=lambda x: str(x.get("sod") or ""), reverse=True)
        ratios: List[float] = []
        for r, ln in zip(rows, ordered):
            y = common.to_number((self.row_snapshot.get(r) or {}).get("yingshou"))
            d = ln.get("deliver")
            if y is None or d is None or float(y) == 0.0:
                return None
            if abs(float(y) - float(d)) > TOL:
                ratios.append(float(d) / float(y))
        kind, ratio = "exact", None
        if ratios:
            ratio = sum(ratios) / len(ratios)
            if not (0.5 < ratio < 1.5):
                return None
            if any(abs(x - ratio) > 5e-4 for x in ratios):
                return None  # 比值不一致 → 更像行错位，不敢认
            kind = "ratio"
        for r, ln in zip(rows, ordered):
            if (ln.get("sod") or "") == sod:
                return r, kind, ratio
        return None

    @staticmethod
    def _is_outstanding(snap: dict) -> bool:
        """拆行后唯一可继续承接回款的行：未结账且没有回款明细。"""
        settled = str(snap.get("jiezhang") or "").strip()
        received = common.to_number(snap.get("huikuan"))
        return settled != "是" and (received is None or abs(float(received)) <= TOL)

    def settled_without_open_row(self, so: str, sod: str = "") -> Optional[int]:
        """已有结账行且同一订单/SOD 不存在未结账行时，返回稳定的已结账行。"""
        so = str(so or "").strip()
        sod = str(sod or "").strip()
        if not so:
            return None
        if sod:
            rows = [
                row_no for row_no in self.sod_index.get(sod, [])
                if str((self.row_snapshot.get(row_no) or {}).get("so") or "").strip() == so
            ]
        else:
            rows = list(self.so_index.get(so, []))
        if not rows:
            return None
        settled_rows = [
            row_no for row_no in sorted(rows)
            if str((self.row_snapshot.get(row_no) or {}).get("jiezhang") or "").strip() == "是"
        ]
        if not settled_rows:
            return None
        if any(
            str((self.row_snapshot.get(row_no) or {}).get("jiezhang") or "").strip() != "是"
            for row_no in rows
        ):
            return None
        return settled_rows[0]

    def so_settlement(self, so: str) -> dict:
        result = settlement_status.inspect_so(so, (
            {"row": row, "so": snap.get("so"), "sod": snap.get("sod"), "settled": snap.get("jiezhang")}
            for row in self.so_index.get(str(so or "").strip(), [])
            for snap in [self.row_snapshot.get(row) or {}]
        ))

        if result['all_settled']:
            incomplete = []
            for sod in result['sods']:
                evidence = [self.row_snapshot[r['row']] for r in result['rows'] if r.get('sod') == sod]
                import receipt_history
                if receipt_history.zero_rows(BR.ledger_rows(self, so, sod)):
                    continue
                if not any((common.to_number(row.get('huikuan')) or 0) > 0
                           and common.norm_date(row.get('shoukuan_time'))
                           and str(row.get('shoukuan_way') or '').strip() for row in evidence):
                    incomplete.append(sod)
            if incomplete:
                result['all_settled'] = False
                result['reason'] = '结账标记为是但缺少完整回款记录，继续逐笔核销：' + ','.join(incomplete)
        return result

    def payment_event_coverage(
        self,
        so: str,
        sod: str,
        amount_local: Optional[float],
        receipt_time: Optional[dt.date],
        payment_way: str,
    ) -> dict:
        """按本批回款的单元格证据判断它是否已经写入盈亏表。

        盈亏表不保存 AR 号和核销记录 NUM，所以已写入证据只能来自同一 SO/SOD
        行的「回款明细 + 收款时间 + 收款方式」。三项必须同时匹配且只能命中一行，
        并且该行已经结账，才能把本批回款判为幂等；整 SO 已结账跳过另由 so_settlement 判断。
        """
        so_s = str(so or "").strip()
        sod_s = str(sod or "").strip()
        if sod_s:
            rows = [
                row_no
                for row_no in self.sod_index.get(sod_s, [])
                if str((self.row_snapshot.get(row_no) or {}).get("so") or "").strip()
                == so_s
            ]
        else:
            rows = list(self.so_index.get(so_s, []))

        expected_date = common.norm_date(receipt_time)
        expected_way = str(payment_way or "").strip()
        expected_amount = common.to_number(amount_local)
        payload = {
            "status": "not_covered",
            "basis": "no_unique_current_event_match",
            "so": so_s,
            "sod": sod_s,
            "candidate_rows": sorted(rows),
            "matched_rows": [],
            "expected": {
                "回款明细": round(float(expected_amount), 2)
                if expected_amount is not None
                else None,
                "收款时间": expected_date,
                "收款方式": expected_way,
            },
        }
        settled_ref = self.settled_without_open_row(so_s, sod_s)
        payload["candidate_values"] = [
            {"row": row_no, "回款明细": common.to_number(snap.get("huikuan")),
             "收款时间": common.norm_date(snap.get("shoukuan_time")),
             "收款方式": str(snap.get("shoukuan_way") or "").strip(),
             "是否结账": str(snap.get("jiezhang") or "").strip()}
            for row_no in sorted(rows) for snap in [self.row_snapshot.get(row_no) or {}]
        ]
        payload["all_rows_settled"] = settled_ref is not None
        payload["settled_row_ref"] = settled_ref

        if expected_amount is None or expected_date is None or not expected_way:
            payload["basis"] = "current_event_signature_incomplete"
            return payload

        matched = []
        for row_no in sorted(rows):
            snap = self.row_snapshot.get(row_no) or {}
            received = common.to_number(snap.get("huikuan"))
            actual_date = common.norm_date(snap.get("shoukuan_time"))
            actual_way = str(snap.get("shoukuan_way") or "").strip()
            if (
                received is not None
                and abs(float(received) - float(expected_amount)) <= TOL
                and actual_date == expected_date
                and actual_way == expected_way
            ):
                matched.append({
                    "row": int(row_no),
                    "settled": str(snap.get("jiezhang") or "").strip() == "是",
                    "回款明细": round(float(received), 2),
                    "收款时间": actual_date,
                    "收款方式": actual_way,
                })
        payload["matched_rows"] = matched
        if len(matched) > 1:
            payload["status"] = "ambiguous"
            payload["basis"] = "current_event_matches_multiple_rows"
        elif len(matched) == 1 and matched[0]["settled"]:
            payload["status"] = "covered"
            payload["basis"] = "ledger_amount_receipt_time_payment_way_and_settled"
            payload["row"] = matched[0]["row"]
        elif len(matched) == 1:
            payload["basis"] = "current_event_row_exists_but_is_not_settled"
        return payload

    def business_rows(self, so: str, sod: str, row: Optional[int] = None) -> List[int]:
        """
        返回同一 SOD 的全部拆分行。

        首次回款前盈亏表可能尚未回填 SOD，此时只能用已唯一命中的当前行作为基线；
        一旦拆行，两行都保留 SOD，后续累计只在该 SOD 内计算，绝不再把同 SO 的其它
        SOD 回款混进来。
        """
        if sod:
            rows = [r for r in self.sod_index.get(sod, [])
                    if str((self.row_snapshot.get(r) or {}).get("so") or "").strip() == so]
            if rows:
                return sorted(rows)
        if row is None:
            return []
        snap = self.row_snapshot.get(int(row)) or {}
        if str(snap.get("so") or "").strip() != so:
            return []
        if sod and str(snap.get("sod") or "").strip() not in ("", sod):
            return []
        return [int(row)]

    def business_totals(
        self, so: str, sod: str, row: Optional[int] = None
    ) -> Tuple[Optional[float], float, List[int]]:
        """返回该 SOD 拆分行的原始应收合计、累计已填回款、行号。"""
        receivable_values: List[float] = []
        received_values: List[float] = []
        rows = self.business_rows(so, sod, row)
        for one_row in rows:
            snap = self.row_snapshot.get(one_row) or {}
            receivable = common.to_number(snap.get("yingshou"))
            received = common.to_number(snap.get("huikuan"))
            if receivable is not None:
                receivable_values.append(float(receivable))
            if received is not None:
                received_values.append(float(received))
        initial_receivable = (
            round(sum(receivable_values), 2) if receivable_values else None
        )
        return initial_receivable, round(sum(received_values), 2), rows

    def comparison_row(
        self,
        so: str,
        sod: str = "",
        current_received: Optional[float] = None,
        receipt_time: Optional[dt.date] = None,
        payment_way: str = "",
    ) -> Optional[int]:
        """
        在离线对比表中定位候选行；仅供差异审计，不参与日常判定。

        拆行后的同一 SOD 可能有多行；优先用“本次回款明细”唯一命中当前切片，
        再用收款时间和收款方式核对当前切片；不能唯一证明就返回 None，禁止猜行。
        """
        candidates: List[int] = []
        if sod:
            candidates = list(self.sod_index.get(sod, []))
        if not candidates and so:
            candidates = list(self.so_index.get(so, []))
        if not candidates:
            return None
        if current_received is not None:
            exact = []
            expected_date = common.norm_date(receipt_time)
            expected_way = str(payment_way or "").strip()
            for one_row in candidates:
                snap = self.row_snapshot.get(one_row) or {}
                got = common.to_number(snap.get("huikuan"))
                if (
                    got is not None
                    and abs(float(got) - float(current_received)) <= TOL
                    and (
                        expected_date is None
                        or common.norm_date(snap.get("shoukuan_time"))
                        == expected_date
                    )
                    and (
                        not expected_way
                        or str(snap.get("shoukuan_way") or "").strip()
                        == expected_way
                    )
                ):
                    exact.append(one_row)
            if len(exact) == 1:
                return exact[0]
        return candidates[0] if len(candidates) == 1 else None

    def so_totals(self, so: str) -> Tuple[Optional[float], float]:
        """兼容旧调用；仅用于没有 SOD 粒度的人工提示。"""
        receivable_values: List[float] = []
        received_values: List[float] = []
        for row in self.so_index.get(so, []):
            snap = self.row_snapshot.get(row) or {}
            receivable = common.to_number(snap.get("yingshou"))
            received = common.to_number(snap.get("huikuan"))
            if receivable is not None:
                receivable_values.append(float(receivable))
            if received is not None:
                received_values.append(float(received))
        return (
            round(sum(receivable_values), 2) if receivable_values else None,
            round(sum(received_values), 2),
        )

    def match(
        self, so: str, sod: str, amount: Optional[float] = None
    ) -> Tuple[Optional[int], str, List[int]]:
        """
        返回 (行号, 依据, 候选)。依据优先级：
          ① SO + 应收金额  ← 主键：回填前后都成立（她一行 = 一个 SOD，应收金额=该SOD交付额）
          ② SOD           ← 她已回填过「实收金额」列时可用（幂等重跑走这条）
          ③ SO 唯一行
        """
        if so and amount is not None:
            rows = self.so_amount_index.get((so, int(round(float(amount) * 100))), [])
            if len(rows) == 1:
                return rows[0], "SO+应收金额", rows
            if len(rows) > 1:
                # 保留 SO+应收金额为主键；只有该键多命中时，才用已经写在盈亏
                # “实收金额”列中的 SOD 在原金额候选集内做二次消歧。SOD 不在
                # 原候选集时不得跨金额选行，避免历史错误 SOD 或金额变化写错行。
                if sod:
                    sod_rows = [
                        r for r in rows
                        if str((self.row_snapshot.get(r) or {}).get("sod") or "").strip()
                        == str(sod).strip()
                    ]
                    if len(sod_rows) == 1:
                        return sod_rows[0], "SO+应收金额+SOD", rows
                    if len(sod_rows) > 1:
                        outstanding = [
                            r for r in sod_rows
                            if self._is_outstanding(self.row_snapshot.get(r) or {})
                        ]
                        if len(outstanding) == 1:
                            return outstanding[0], "SO+应收金额+SOD未结清行", rows
                return None, "E8", rows
        if sod:
            rows = self.sod_index.get(sod, [])
            if len(rows) == 1:
                return rows[0], "SOD", rows
            if len(rows) > 1:
                outstanding = [
                    r for r in rows if self._is_outstanding(self.row_snapshot.get(r) or {})
                ]
                if len(outstanding) == 1:
                    return outstanding[0], "SOD未结清行", rows
                return None, "E8", rows
        if so:
            rows = self.so_index.get(so, [])
            if len(rows) == 1:
                return rows[0], "SO", rows
            if len(rows) > 1:
                return None, "E8", rows
            return None, "E2", []
        return None, "E7", []
